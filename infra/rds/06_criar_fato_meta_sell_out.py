"""Cria e carrega estoque_redes.fato_meta_sell_out a partir de "meta sell out.xlsx".

Meta mensal de faturamento de sell out por SKU (2026). Roda como rubens_dba pelo
túnel (infra/rds/abrir_tunel_wsl.sh), numa transação só: cria a tabela, carrega,
dá as mesmas permissões de leitura da estoque_redes.fato_sell_in e confere. Se a
tabela já existir, para sem mexer em nada. Senha lida do bi/.env, nunca impressa.

Pré-requisito (regra 7 da Fase 9): snapshot manual do RDS antes.

    uv run --with openpyxl python infra/rds/06_criar_fato_meta_sell_out.py "meta sell out.xlsx"
"""
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import psycopg2
from dotenv import dotenv_values
from openpyxl import load_workbook

ARQUIVO = Path(sys.argv[1] if len(sys.argv) > 1 else "meta sell out.xlsx")
ENV = Path(__file__).resolve().parents[2] / ".env"

linhas = list(load_workbook(ARQUIVO, data_only=True, read_only=True).worksheets[0].iter_rows(values_only=True))
assert tuple(linhas[0][:4]) == ("data", "sku", "cod_apres", "meta_faturamento"), linhas[0]
dados = [(r[0].date(), str(r[1]).strip(), int(r[2]), Decimal(str(r[3])).quantize(Decimal("0.01"), ROUND_HALF_UP))
         for r in linhas[1:] if any(v is not None for v in r)]
total_planilha = sum(d[3] for d in dados)

con = psycopg2.connect(host="127.0.0.1", port=15432, dbname="easelabs", user="rubens_dba",
                       password=dotenv_values(ENV).get("COCKPIT_DB_PASS"), sslmode="require", connect_timeout=15)
try:
    with con, con.cursor() as cur:
        cur.execute("SELECT to_regclass('estoque_redes.fato_meta_sell_out')")
        if cur.fetchone()[0]:
            raise SystemExit("estoque_redes.fato_meta_sell_out já existe: nada foi alterado.")
        cur.execute("""
            CREATE TABLE estoque_redes.fato_meta_sell_out (
                data              date          NOT NULL,
                sku               text          NOT NULL,
                cod_apres         bigint        NOT NULL,
                meta_faturamento  numeric(15,2) NOT NULL,
                PRIMARY KEY (data, cod_apres)
            )""")
        cur.execute("""COMMENT ON TABLE estoque_redes.fato_meta_sell_out IS
            'Meta mensal de faturamento (R$) de sell out da Ease por SKU. data = 1º dia do mês; '
            'cod_apres liga com cddd.vw_sell_out.cod_apres. Carregada de "meta sell out.xlsx" em 2026-10-14.'""")
        cur.executemany("INSERT INTO estoque_redes.fato_meta_sell_out VALUES (%s, %s, %s, %s)", dados)
        # Mesma leitura da fato_sell_in: aplicações, BI, Jarvis (bi_chatbot_ro) e os DBAs.
        for papel in ("app_pipelines", "app_api", "app_eventos", "bi_leitura", "dbt_transform",
                      "bi_chatbot_ro", "app_forecast_estoque"):
            cur.execute(f"GRANT SELECT ON estoque_redes.fato_meta_sell_out TO {papel}")
        cur.execute("GRANT ALL ON estoque_redes.fato_meta_sell_out TO paulo_dba")
        cur.execute("SELECT count(*), sum(meta_faturamento), min(data), max(data) FROM estoque_redes.fato_meta_sell_out")
        n, soma, ini, fim = cur.fetchone()
        assert n == len(dados) and soma == total_planilha, (n, soma, total_planilha)
    print(f"ok: {n} linhas, total R$ {soma:,.2f}, de {ini} a {fim}")
finally:
    con.close()
