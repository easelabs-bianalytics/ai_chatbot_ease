"""Conferência da planilha antes de devolver (ADR-0031).

Determinística, sem modelo: compara o arquivo que a pessoa mandou com o que
o Jarvis vai devolver e barra a entrega se algo mudou fora do pedido.

- toda aba original continua lá, com o mesmo nome;
- nenhuma linha sumiu ou entrou na base;
- nenhuma célula original com valor (ou fórmula) mudou, a não ser nas
  colunas que o pedido mandou reescrever;
- célula vazia só ganhou valor nas colunas de destino;
- aba nova só as pedidas e a "Notas do Jarvis".

E mede a cobertura de cada coluna escrita: quantas linhas ganharam valor e
qual valor domina, quando um domina ("SEM REP" em 146 das 230, conversa 44).
É com esses números que a resposta fala da planilha, sem contradizê-la.
"""

import io
import math
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from attachments.planilha import (
    NOME_DAS_NOTAS,
    _linhas_do_csv,
    _vazio,
    indice_do_cabecalho,
    normalizar,
)

# A partir de quanto um valor "domina" a coluna e merece ser dito.
PARTE_DOMINANTE = 0.5
MAX_EXEMPLOS_DE_PROBLEMA = 5


@dataclass(frozen=True)
class Cobertura:
    aba: str
    coluna: str
    com_valor: int
    total: int
    dominante: str = ""
    vezes_do_dominante: int = 0


@dataclass
class Conferencia:
    problemas: list = field(default_factory=list)
    cobertura: list = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problemas


def _abas(nome: str, dados: bytes) -> dict:
    """{título: linhas} com fórmulas como texto. CSV vira uma aba só; se o
    arquivo final for xlsx vindo de CSV, a base é a aba "Dados"."""
    if Path(nome).suffix.lower() == ".csv":
        return {"": [list(l) for l in _linhas_do_csv(dados)]}
    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=False)
    try:
        return {aba.title: [list(l) for l in aba.iter_rows(values_only=True)] for aba in livro.worksheets}
    finally:
        livro.close()


# Tolerância relativa para número. O openpyxl grava com 16 algarismos o que
# veio com 17 (167.26185324474665 → 167.2618532447466): o Excel guarda 15,
# então não é mudança nenhuma — e barrou a primeira aba nova da planilha de
# metas (2026-10-01). Qualquer alteração de verdade fica muito acima disto.
TOLERANCIA_RELATIVA = 1e-12


def _mesmo(a, b) -> bool:
    if _vazio(a) and _vazio(b):
        return True
    numeros = (int, float)
    if isinstance(a, numeros) and isinstance(b, numeros) and not isinstance(a, bool) and not isinstance(b, bool):
        return math.isclose(a, b, rel_tol=TOLERANCIA_RELATIVA, abs_tol=0.0)
    if isinstance(a, str) or isinstance(b, str):
        # CSV lido de novo é texto; o número escrito no xlsx vindo de CSV
        # também foi texto. Comparar a forma escrita basta.
        return str(a).strip() == str(b).strip()
    return a == b


def _celula(linhas, i, j):
    if i < len(linhas) and j < len(linhas[i]):
        return linhas[i][j]
    return None


def conferir(nome_original: str, original: bytes, nome_final: str, final: bytes, permissoes) -> Conferencia:
    """`permissoes`: {aba (título; "" = primeira): {"escrever": {cabeçalhos},
    "sobrescrever": {cabeçalhos}}} e, na chave especial `"_abas_novas"`, os
    títulos das abas criadas."""
    conferencia = Conferencia()
    antes = _abas(nome_original, original)
    depois = _abas(nome_final, final)
    titulos_antes = list(antes)
    titulos_depois = list(depois)
    de_csv = Path(nome_original).suffix.lower() == ".csv" and Path(nome_final).suffix.lower() != ".csv"
    if de_csv:
        # O CSV virou xlsx para caber aba nova: a base é a aba "Dados".
        mapa = {"": "Dados"}
    else:
        mapa = {t: t for t in titulos_antes}
    novas = {normalizar(t) for t in permissoes.get("_abas_novas", ())} | {normalizar(NOME_DAS_NOTAS)}

    for titulo in titulos_depois:
        if titulo not in mapa.values() and normalizar(titulo) not in novas:
            conferencia.problemas.append(f'apareceu a aba "{titulo}", que ninguém pediu')

    primeira = titulos_antes[0] if titulos_antes else ""
    for titulo in titulos_antes:
        destino = mapa.get(titulo)
        if destino not in depois:
            conferencia.problemas.append(f'a aba "{titulo or "Dados"}" sumiu')
            continue
        regra = permissoes.get(titulo) or (permissoes.get("") if titulo == primeira else None) or {}
        escrever = {normalizar(c) for c in regra.get("escrever", ())}
        sobrescrever = {normalizar(c) for c in regra.get("sobrescrever", ())}
        _conferir_aba(conferencia, titulo or "Dados", antes[titulo], depois[destino], escrever, sobrescrever)
    return conferencia


def _conferir_aba(conferencia, titulo, antes, depois, escrever, sobrescrever) -> None:
    i_cabecalho = indice_do_cabecalho(depois)
    if i_cabecalho is None:
        if any(any(not _vazio(v) for v in linha) for linha in antes):
            conferencia.problemas.append(f'a aba "{titulo}" ficou vazia')
        return
    cabecalho = [normalizar(v) for v in depois[i_cabecalho]]
    ultima_antes = max((i for i, l in enumerate(antes) if any(not _vazio(v) for v in l)), default=-1)
    ultima_depois = max((i for i, l in enumerate(depois) if any(not _vazio(v) for v in l)), default=-1)
    if ultima_depois != ultima_antes:
        conferencia.problemas.append(
            f'a aba "{titulo}" tinha {ultima_antes + 1} linhas e ficou com {ultima_depois + 1}'
        )

    largura = max([len(l) for l in antes + depois] or [0])
    mudadas = []
    for i in range(max(len(antes), len(depois))):
        for j in range(largura):
            a, b = _celula(antes, i, j), _celula(depois, i, j)
            if _mesmo(a, b):
                continue
            coluna = cabecalho[j] if j < len(cabecalho) else ""
            if i == i_cabecalho and _vazio(a) and coluna in escrever:
                continue  # o cabeçalho da coluna nova
            if i > i_cabecalho and _vazio(a) and coluna in escrever:
                continue  # célula vazia que o pedido mandou preencher
            if i > i_cabecalho and not _vazio(a) and coluna in sobrescrever:
                continue  # reescrita pedida explicitamente
            mudadas.append(f"linha {i + 1}, coluna {j + 1}")
    if mudadas:
        exemplos = ", ".join(mudadas[:MAX_EXEMPLOS_DE_PROBLEMA])
        conferencia.problemas.append(
            f'{len(mudadas)} células da aba "{titulo}" mudaram fora do pedido ({exemplos})'
        )

    corpo = [l for l in depois[i_cabecalho + 1:] if any(not _vazio(v) for v in l)]
    for j, coluna in enumerate(depois[i_cabecalho]):
        if normalizar(coluna) not in escrever:
            continue
        valores = [l[j] for l in corpo if j < len(l) and not _vazio(l[j])]
        contagem = Counter(str(v).strip() for v in valores)
        dominante, vezes = contagem.most_common(1)[0] if contagem else ("", 0)
        if vezes < max(3, PARTE_DOMINANTE * len(valores)):
            dominante, vezes = "", 0
        conferencia.cobertura.append(Cobertura(
            aba=titulo, coluna=str(coluna), com_valor=len(valores), total=len(corpo),
            dominante=dominante, vezes_do_dominante=vezes,
        ))

