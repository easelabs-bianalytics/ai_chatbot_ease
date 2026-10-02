"""A Base 660: retrato fixo de médicos que o Marketing trabalha (ADR-0032).

Diferente da Área Médica e do Email MKT, não há API: a base é uma planilha
("Médicos 660 - Informações MKT.xlsx") carregada uma vez em
`marketing.medicos_660` pelo comando `carregar_medicos_660`. Carregar de novo
substitui tudo, numa transação, como as outras fontes.

O `crm_link` sai da mesma função das outras bases (`marketing.crm`): a
planilha traz o CRM LINK sem os zeros (`PI4896`), e o BI usa UF + 7 dígitos
(`PI0004896`). O que veio fica em `crm_link_planilha`, para conferência.

As colunas de presença da planilha (Área Médica, Email MKT, já prescreveu,
cobertura e painel da força de vendas) não são carregadas: eram o retrato do
dia em que a planilha foi feita, e o Jarvis cruza isso na hora, com o dado
de hoje, pelo `crm_link`.
"""

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from psycopg2.extras import execute_values

from marketing.crm import crm_link

FONTE = "medicos_660"

# Título da coluna na planilha → coluna da tabela.
COLUNAS_DA_PLANILHA = {
    "Nome": "nome",
    "Número Registro": "crm_numero",
    "UF Registro": "uf_registro",
    "Especialidade Principal": "especialidade",
    "Especialidade 2": "especialidade_2",
    "Estado": "estado",
    "Cidade": "cidade",
    "CEP": "cep",
    "Bairro": "bairro",
    "Endereço": "endereco",
    "Telefone": "telefone",
    "Telefone 2": "telefone_2",
    "Telefone 3": "telefone_3",
    "E-Mail": "email",
    "Valor Online": "valor_online",
    "Valor Presencial": "valor_presencial",
    "CRM LINK": "crm_link_planilha",
}
COLUNAS = (*COLUNAS_DA_PLANILHA.values(), "crm_link", "carregado_em")
_VALORES = ("valor_online", "valor_presencial")


def _texto(valor):
    """Texto limpo; "-" e vazio são ausência."""
    texto = " ".join(str(valor).split()) if valor is not None else ""
    return None if texto in ("", "-") else texto


def _valor(valor):
    """Preço da consulta, ou None ("Indisponível", "-")."""
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    try:
        return Decimal(str(valor).replace(",", ".").strip())
    except (InvalidOperation, AttributeError):
        return None


def linhas_da_planilha(linhas, agora=None) -> list:
    """Da planilha (cabeçalho + linhas, como o openpyxl devolve) às tuplas
    na ordem de `COLUNAS`. Sem a coluna obrigatória, falha antes de gravar."""
    agora = agora or datetime.now(timezone.utc)
    cabecalho, *dados = list(linhas)
    posicao = {_texto(titulo): i for i, titulo in enumerate(cabecalho)}
    faltando = [titulo for titulo in COLUNAS_DA_PLANILHA if titulo not in posicao]
    if faltando:
        raise ValueError(f"colunas ausentes na planilha da Base 660: {', '.join(faltando)}")

    saida = []
    for linha in dados:
        if not any(v not in (None, "") for v in linha):
            continue
        registro = {}
        for titulo, coluna in COLUNAS_DA_PLANILHA.items():
            bruto = linha[posicao[titulo]] if posicao[titulo] < len(linha) else None
            registro[coluna] = _valor(bruto) if coluna in _VALORES else _texto(bruto)
        registro["crm_link"] = crm_link(registro["crm_numero"], registro["uf_registro"]) or None
        registro["carregado_em"] = agora
        saida.append(tuple(registro[c] for c in COLUNAS))
    return saida


def ler_planilha(caminho) -> list:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        return linhas_da_planilha(livro.worksheets[0].iter_rows(values_only=True))
    finally:
        livro.close()


def gravar_base_660(conexao, linhas) -> dict:
    with conexao.cursor() as cursor:
        cursor.execute("DELETE FROM marketing.medicos_660")
        execute_values(cursor, f"INSERT INTO marketing.medicos_660 ({', '.join(COLUNAS)}) VALUES %s",
                       linhas, page_size=1000)
    indice = COLUNAS.index("crm_link")
    links = {l[indice] for l in linhas if l[indice]}
    return {"medicos": len(linhas), "com_crm_link": sum(1 for l in linhas if l[indice]), "crm_links": len(links)}
