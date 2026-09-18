"""Planilha Excel com o resultado de uma consulta (ADR-0020).

A planilha sai direto do banco, sem passar pela IA: não custa token e não
tem como trazer número inventado. O que ela precisa é parecer um relatório
do time de BI — cabeçalho da marca, bordas, filtro, primeira linha
congelada, números e datas no formato certo — e dizer de onde veio, numa
aba de informações.
"""

import datetime
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INDIGO = "5558D4"
CINZA_BORDA = "DDDFE6"
CINZA_ZEBRA = "F7F8FA"
CINZA_TEXTO = "5E6484"

# Código não é número: CNPJ, EAN e CRM perdem o zero à esquerda e viram
# notação científica se o Excel os tratar como número.
_COLUNA_DE_CODIGO = re.compile(r"(cnpj|cpf|ean|crm|cep|cod|telefone|telef|ddd|setor|anomes)", re.I)
_DATA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATA_HORA = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")

_BORDA = Border(*(Side(style="thin", color=CINZA_BORDA),) * 4)
_FONTE = "Calibri"


def _titulo_da_coluna(nome: str) -> str:
    """`total_und` → `Total und`: a planilha vai para quem não lê SQL."""
    texto = str(nome).replace("_", " ").strip()
    return texto[:1].upper() + texto[1:] if texto else texto


def _valor(coluna: str, valor):
    if valor is None:
        return None
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if _COLUNA_DE_CODIGO.search(coluna):
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        return str(valor)
    if isinstance(valor, str):
        if _DATA.match(valor):
            try:
                return datetime.date.fromisoformat(valor)
            except ValueError:
                return valor
        if _DATA_HORA.match(valor):
            try:
                # O Excel não guarda fuso: fica a hora local de Brasília.
                return datetime.datetime.fromisoformat(valor).replace(tzinfo=None)
            except ValueError:
                return valor
    return valor


def _formato(valor) -> str | None:
    if isinstance(valor, datetime.datetime):
        return "dd/mm/yyyy hh:mm"
    if isinstance(valor, datetime.date):
        return "dd/mm/yyyy"
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return "#,##0"
    if isinstance(valor, float):
        return "#,##0.00" if not valor.is_integer() else "#,##0"
    return None


def montar_planilha(colunas, linhas, info: dict) -> bytes:
    """Devolve o `.xlsx` pronto.

    `info` vai para a aba "Informações": pergunta, momento, linhas, se foi
    cortado, referência e a consulta executada."""
    livro = Workbook()
    dados = livro.active
    dados.title = "Dados"

    cabecalho = [_titulo_da_coluna(c) for c in colunas]
    dados.append(cabecalho)
    for celula in dados[1]:
        celula.font = Font(name=_FONTE, bold=True, color="FFFFFF", size=11)
        celula.fill = PatternFill("solid", fgColor=INDIGO)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = _BORDA
    dados.row_dimensions[1].height = 22

    larguras = [len(t) for t in cabecalho]
    zebra = PatternFill("solid", fgColor=CINZA_ZEBRA)
    for n, linha in enumerate(linhas, start=2):
        valores = [_valor(colunas[i], v) for i, v in enumerate(linha)]
        dados.append(valores)
        for i, valor in enumerate(valores):
            celula = dados.cell(row=n, column=i + 1)
            celula.border = _BORDA
            celula.font = Font(name=_FONTE, size=10)
            formato = _formato(valor)
            if formato:
                celula.number_format = formato
            if n % 2 == 1:
                celula.fill = zebra
            texto = valor.strftime("%d/%m/%Y") if isinstance(valor, datetime.date) else str(valor or "")
            larguras[i] = max(larguras[i], len(texto))

    for i, largura in enumerate(larguras, start=1):
        dados.column_dimensions[get_column_letter(i)].width = min(max(largura + 3, 10), 60)
    dados.freeze_panes = "A2"
    if colunas:
        dados.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(len(linhas) + 1, 1)}"

    _aba_de_informacoes(livro, info)

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _aba_de_informacoes(livro, info: dict) -> None:
    aba = livro.create_sheet("Informações")
    aba.column_dimensions["A"].width = 24
    aba.column_dimensions["B"].width = 100

    aba["A1"] = "Jarvis · Ease Labs"
    aba["A1"].font = Font(name=_FONTE, bold=True, size=14, color=INDIGO)
    aba["A2"] = "BI & Analytics — dados consultados em modo somente leitura na base de BI"
    aba["A2"].font = Font(name=_FONTE, size=10, color=CINZA_TEXTO)

    campos = [
        ("Pergunta", info.get("pergunta", "")),
        ("Planilha gerada em", info.get("gerada_em", "")),
        ("Linhas", info.get("linhas", "")),
        ("Observação", info.get("observacao", "")),
        ("Referência usada", info.get("referencia") or "nenhuma (consulta escrita pela IA)"),
        ("Consulta executada", info.get("sql", "")),
    ]
    for n, (rotulo, valor) in enumerate(campos, start=4):
        a = aba.cell(row=n, column=1, value=rotulo)
        b = aba.cell(row=n, column=2, value=valor)
        a.font = Font(name=_FONTE, bold=True, size=10)
        b.font = Font(name="Consolas" if rotulo == "Consulta executada" else _FONTE, size=10)
        a.alignment = Alignment(vertical="top")
        b.alignment = Alignment(vertical="top", wrap_text=True)
        a.border = b.border = _BORDA
    aba.row_dimensions[9].height = min(15 * (str(info.get("sql", "")).count("\n") + 1), 400)
