"""Nenhum limite corta a resposta em silêncio (2026-09-25).

Depois do "IC" cortado no limite de tokens e da planilha do WhatsApp presa
em 100 linhas (conversa do Paulo: 340 médicos, planilha com 100), o Rubens
pediu: nenhum limite pode tirar dado da resposta sem que a pessoa saiba.
Onde dá, entrega tudo — a planilha de cada tabela refaz a consulta DELA; onde
um teto é necessário, a resposta diz o que ficou de fora.
"""

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from ai_orchestrator import orchestrator
from ai_orchestrator.providers.openai_provider import _resposta_incompleta
from attachments.planilha import ler_estrutura
from datasource.executors.fake import FakeQueryExecutor, make_result
from messaging import planilha_da_resposta
from messaging.models import Message
from tests.fakes.providers import ScriptedAIProvider, plano, resposta
from tests.unit.test_orchestrator import _pergunta, _responder, catalogo, conversa  # noqa: F401
from whatsapp import graficos, saida

pytestmark = pytest.mark.django_db

EVOLUCAO = make_result(("competencia", "px"), [("2026-01-01", 4327.0), ("2026-02-01", 4821.0)])
RANKING = make_result(("territorio", "medico", "px"), [(f"T{i // 10}", f"MEDICO {i}", 400.0 - i) for i in range(340)])
SQL_EVOLUCAO = "SELECT data AS competencia, SUM(und) AS px FROM cddd.vendas_consolidado GROUP BY 1"
SQL_RANKING = "SELECT canal AS territorio, canal AS medico, SUM(und) AS px FROM cddd.vendas_consolidado GROUP BY 1"


def _entregas(conversa, catalogo, executor, texto="Neurologia lidera com 400 PX.", consultas=None):
    consultas = consultas or (
        {"titulo": "Evolução", "sql": SQL_EVOLUCAO, "reference_query_id": "A01"},
        {"titulo": "Ranking por território", "sql": SQL_RANKING, "reference_query_id": "A06"},
    )
    provider = ScriptedAIProvider(
        [plano("", consultas=tuple(consultas), entendimento="evolução e ranking")],
        [resposta(texto, blocos=(
            {"tipo": "texto", "texto": texto},
            {"tipo": "tabela", "consulta": 1, "colunas": ["territorio", "medico", "px"]},
        ))],
    )
    reply = _responder(_pergunta(conversa, "a evolução e o ranking dos prescritores"), catalogo,
                       provider=provider, executor=executor)
    return reply, provider


def test_itens_1_e_2_planilha_de_resposta_com_varias_consultas_vem_inteira_e_da_tabela_certa(conversa, catalogo):
    """Com várias consultas, a planilha do WhatsApp saía com as 100 linhas
    guardadas para a tela — ou refazia "a última consulta", que era de outra
    tabela. Agora cada tabela guarda a SUA consulta, e a planilha refaz essa."""
    reply, _ = _entregas(conversa, catalogo, FakeQueryExecutor([EVOLUCAO, RANKING]))
    dados = reply.raw_response["dados_blocos"]["1"]
    assert dados["sql"] == SQL_RANKING and dados["total"] == 340 and len(dados["rows"]) == 100

    saida_msg = Message.objects.get(in_reply_to=reply.message, direction=Message.Direction.OUTBOUND)
    executor = FakeQueryExecutor([RANKING])
    envios = saida.montar(saida_msg, executor=executor)

    assert "vendas_consolidado" in executor.executed[0] and "territorio" in executor.executed[0]
    documento = next(e for e in envios if e.tipo == "documento")
    linhas = max(aba.max_row for aba in load_workbook(io.BytesIO(documento.dados)).worksheets)
    assert linhas >= 341
    texto = " ".join(e.texto for e in envios if e.tipo == "texto")
    assert "e mais 335 linhas" in texto


def test_item_1_se_a_consulta_nao_refaz_a_legenda_diz_o_total_real(conversa, catalogo):
    from datasource.executors.base import QueryTimeout

    reply, _ = _entregas(conversa, catalogo, FakeQueryExecutor([EVOLUCAO, RANKING]))
    saida_msg = Message.objects.get(in_reply_to=reply.message, direction=Message.Direction.OUTBOUND)

    envios = saida.montar(saida_msg, executor=FakeQueryExecutor([QueryTimeout("demorou")]))

    documento = next(e for e in envios if e.tipo == "documento")
    assert "100 de 340" in documento.texto
    assert "primeiras 100 de 340" in " ".join(e.texto for e in envios if e.tipo == "texto")


def test_item_2_botao_do_site_baixa_a_planilha_da_tabela_pedida(conversa, catalogo):
    reply, _ = _entregas(conversa, catalogo, FakeQueryExecutor([EVOLUCAO, RANKING]))
    saida_msg = Message.objects.get(in_reply_to=reply.message, direction=Message.Direction.OUTBOUND)
    executor = FakeQueryExecutor([RANKING])

    planilha = planilha_da_resposta.gerar(saida_msg, conversa.user, executor, consulta=1)

    assert planilha.linhas == 340 and not planilha.erro
    assert "territorio" in executor.executed[0]


def test_item_3_consulta_cortada_no_limite_e_avisada_a_redacao_e_marcada_na_tabela(conversa, catalogo):
    cortado = make_result(RANKING.columns, RANKING.rows, truncated=True)
    reply, provider = _entregas(conversa, catalogo, FakeQueryExecutor([EVOLUCAO, cortado]))

    assert provider.answer_requests[0].consultas[1]["truncated"] is True
    assert reply.raw_response["dados_blocos"]["1"]["truncado"] is True


def test_item_4_tabela_de_seguranca_diz_quantas_mostra():
    texto = orchestrator._tabela(make_result(("rede", "und"), [(f"R{i}", i) for i in range(45)]))

    assert "Mostrando as 20 primeiras de 45 linhas" in texto


def test_item_5_pizza_do_whatsapp_soma_o_resto_em_outras():
    """Mostrava as 8 maiores e descartava o resto: proporções erradas."""
    dados = {"columns": ["rede", "und"], "rows": [[f"R{i}", 100 - i] for i in range(12)]}

    spec = graficos.de_simples({"tipo": "pizza", "x": "rede", "series": ["und"], "titulo": "Share"}, dados)

    nomes = [v["rede"] for v in spec["data"]["values"]]
    assert len(nomes) == 9 and nomes[-1] == "Outras"
    assert spec["data"]["values"][-1]["und"] == sum(100 - i for i in range(8, 12))
    assert "Outras" in spec["title"]["subtitle"]


def test_item_5_barras_do_whatsapp_avisam_o_corte():
    dados = {"columns": ["rede", "und"], "rows": [[f"R{i}", 100 - i] for i in range(20)]}

    spec = graficos.de_simples({"tipo": "barras", "x": "rede", "series": ["und"], "titulo": ""}, dados)

    assert "Mostrando 14 de 20" in spec["title"]["subtitle"]


def test_item_7_texto_nao_promete_planilha_que_nao_foi():
    texto = "A lista está no botão “Baixar Excel” logo abaixo desta resposta."

    assert "na planilha anexada" in saida._sem_o_botao_da_tela(texto, tem_planilha=True)
    sem = saida._sem_o_botao_da_tela(texto, tem_planilha=False)
    assert "anexada" not in sem and "Não consegui gerar a planilha" in sem


def test_item_11_resposta_marcada_incompleta_e_reconhecida_como_cortada():
    """Sem isto a resposta cortada virava "falha da IA", e a segunda
    tentativa mais curta nem acontecia."""
    incompleta = SimpleNamespace(status="incomplete", incomplete_details=SimpleNamespace(reason="max_output_tokens"))
    completa = SimpleNamespace(status="completed", incomplete_details=None)

    assert _resposta_incompleta(incompleta) and not _resposta_incompleta(completa)


def test_item_13_series_que_nao_cabem_ficam_registradas_para_o_aviso(conversa, catalogo):
    resultado = make_result(("mes", "a", "b", "c", "d"), [("2026-01", 1, 2, 3, 4), ("2026-02", 2, 3, 4, 5)])
    mensagem = SimpleNamespace(content="evolução")

    grafico = orchestrator._grafico({"tipo": "linha", "x": "mes", "series": ["a", "b", "c", "d"], "titulo": ""},
                                    resultado, mensagem, SimpleNamespace(sql="select 1"))

    assert grafico["series"] == ["a", "b", "c"] and grafico["series_de_fora"] == ["d"]


def test_item_14_entregas_alem_do_teto_sao_ditas(conversa, catalogo):
    consultas = tuple({"titulo": f"Parte {i}", "sql": SQL_EVOLUCAO, "reference_query_id": "A01"} for i in range(1, 7))
    provider = ScriptedAIProvider([plano("", consultas=consultas, entendimento="seis partes")],
                                  [resposta("Foram 4327 PX em jan/2026.")])

    _responder(_pergunta(conversa, "seis coisas"), catalogo, provider=provider,
               executor=FakeQueryExecutor([EVOLUCAO] * 4))

    nota = provider.answer_requests[0].pedido_nao_atendido
    assert "6 partes" in nota and "«Parte 5»" in nota and "«Parte 6»" in nota


def test_item_15_planilha_com_colunas_demais_avisa_no_resumo():
    cabecalho = ";".join(f"col{i}" for i in range(65))
    linha = ";".join(str(i) for i in range(65))
    estrutura = ler_estrutura("larga.csv", f"{cabecalho}\n{linha}\n".encode())

    assert "65 colunas" in estrutura.resumo and "Diga isso na resposta" in estrutura.resumo


def test_item_6_mensagem_longa_no_whatsapp_e_recusada_com_aviso(monkeypatch):
    from whatsapp import entrada, servicos
    from whatsapp.cliente import ClienteFake

    recebida = entrada.Recebida(tipo=entrada.TEXTO, texto="x" * (entrada.MAX_CARACTERES + 1), jid="5511@s.whatsapp.net",
                                id="L1", numero="5511977776666", grupo=False, nome="Paulo")
    monkeypatch.setattr(entrada, "ler", lambda payload: recebida)
    monkeypatch.setattr(servicos, "_dono", lambda r: (SimpleNamespace(pk=1), None))
    cliente = ClienteFake()

    assert servicos.receber({}, cliente=cliente) == "mensagem_longa"
    assert "8.000 caracteres" in cliente.enviados[-1]["texto"]
