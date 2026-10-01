"""A planilha como tabela da consulta, num Postgres de verdade (ADR-0031).

Roda no `analytics_db` sintético, com o usuário somente leitura — o mesmo
desenho do RDS. Prova o que o fake não prova: que a CTE com `unnest` e os
valores como parâmetro é SQL válido numa transação READ ONLY, que o `%` do
modelo sobrevive aos parâmetros e que o JOIN com o cadastro traz uma linha
por linha da planilha, inclusive a que não casou.
"""

import io
import os

import pytest
from openpyxl import Workbook, load_workbook

from attachments import anexo_sql, qa
from attachments.planilha import PedidoDePreenchimento, ler_estrutura, preencher
from catalog.loader import load_catalog
from datasource.executors.postgres_readonly import PostgresReadOnlyExecutor
from datasource.sql_guard import validate_sql

DSN_PADRAO = "postgresql://bi_readonly:bi_readonly_dev@localhost:5435/analytics"


def _xlsx(linhas) -> bytes:
    livro = Workbook()
    livro.active.title = "Painel"
    for linha in linhas:
        livro.active.append(list(linha))
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


# CRMs escritos como a pessoa escreve: minúsculo, com traço, e um que não
# existe no cadastro. No banco sintético: MG0000001 → UTC 101, SP0000002 → 202.
PAINEL = _xlsx([
    ("Painel de teste",),
    ("Nome", "CRM", "Obs"),
    ("Ana", "mg-0000001", "50% do painel"),
    ("Bia", "SP0000002", None),
    ("Caio", "RJ9999999", "O'Neil; DROP TABLE x"),
])


@pytest.fixture(scope="module")
def executor():
    return PostgresReadOnlyExecutor(
        dsn=os.environ.get("TEST_ANALYTICS_DATABASE_URL", DSN_PADRAO), statement_timeout_ms=5000
    )


@pytest.fixture(scope="module")
def tabelas():
    return anexo_sql.tabelas(ler_estrutura("painel.xlsx", PAINEL))


SQL = """
SELECT a._linha, a.crm, btrim(fv.desc_territorio) AS representante
FROM anexo.painel a
LEFT JOIN audit.medico m
       ON upper(regexp_replace(m.crm, '[^A-Za-z0-9]', '', 'g'))
        = upper(regexp_replace(a.crm, '[^A-Za-z0-9]', '', 'g'))
LEFT JOIN (SELECT DISTINCT cod_utc, desc_territorio FROM cddd.forca_vendas) fv ON fv.cod_utc = m.utc_codigo
WHERE a.obs IS NULL OR a.obs NOT LIKE '%nada%'
ORDER BY a._linha
"""


def _rodar(executor, tabelas, sql, max_rows=500):
    guard = validate_sql(sql, load_catalog(), max_rows=max_rows, anexo=set(tabelas))
    assert guard.approved, guard.reason
    return anexo_sql.ExecutorComAnexo(executor, tabelas).run(guard.sql, max_rows=max_rows)


def test_join_com_o_cadastro_traz_uma_linha_por_linha_da_planilha(executor, tabelas):
    resultado = _rodar(executor, tabelas, SQL)

    assert resultado.columns == ("_linha", "crm", "representante")
    assert resultado.rows == (
        (3, "mg-0000001", "TERRITORIO MG CENTRO"),
        (4, "SP0000002", "TERRITORIO SP CAPITAL"),
        (5, "RJ9999999", None),
    )


def test_valor_de_celula_com_aspas_e_so_dado(executor, tabelas):
    resultado = _rodar(executor, tabelas, "SELECT a.obs FROM anexo.painel a WHERE a._linha = 5")

    assert resultado.rows == (("O'Neil; DROP TABLE x",),)


def test_analise_agregada_sobre_a_planilha(executor, tabelas):
    resultado = _rodar(executor, tabelas, "SELECT count(*) AS medicos, count(a.obs) AS com_obs FROM anexo.painel a")

    assert resultado.rows == ((3, 2),)


def test_do_banco_ao_arquivo_conferido(executor, tabelas):
    """O caminho inteiro da conversa 44, com banco de verdade: consulta sobre
    a planilha, preenchimento pela linha e conferência."""
    resultado = _rodar(executor, tabelas, SQL)
    pedido = PedidoDePreenchimento.do_plano({
        "coluna_chave": "CRM", "chave_no_resultado": "_linha", "aba": "",
        "colunas": [{"coluna_destino": "Representante", "valor_no_resultado": "representante"}],
    })

    feita = preencher("painel.xlsx", PAINEL, pedido, resultado.columns, resultado.rows)
    conferencia = qa.conferir("painel.xlsx", PAINEL, "painel.xlsx", feita.dados,
                              {"Painel": {"escrever": {"Representante"}, "sobrescrever": set()}})

    assert conferencia.ok, conferencia.problemas
    # O LEFT JOIN traz a linha do CRM que não existe, com o valor vazio.
    assert feita.vazias_nomes == ("RJ9999999",)
    valores = [l[3] for l in load_workbook(io.BytesIO(feita.dados))["Painel"].iter_rows(min_row=3, values_only=True)]
    assert valores == ["TERRITORIO MG CENTRO", "TERRITORIO SP CAPITAL", None]
