"""Motor de anexos do Jarvis 3.0 (ADR-0031), peça por peça, sem banco.

O que estes testes fixam:

1. o perfil acha o cabeçalho abaixo do título, nomeia as colunas para o SQL,
   aponta a chave e as colunas calculadas — e guarda as linhas com o número
   delas no arquivo;
2. `anexo.<aba>` vira CTE com os valores como PARÂMETRO, nunca no texto;
3. o validador só aceita as abas da planilha da conversa;
4. o preenchimento pela `_linha` põe cada valor na linha de onde veio, não
   reescreve o que a pessoa já tinha e não conta linha vazia como feita;
5. a conferência barra qualquer mudança fora do pedido;
6. aba nova e "Notas do Jarvis" entram sem tocar na base.
"""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from attachments import anexo_sql, qa
from attachments.limites import AnexoRecusado
from attachments.planilha import (
    NOME_DAS_NOTAS,
    PedidoDePreenchimento,
    adicionar_aba,
    anotar,
    indice_do_cabecalho,
    ler_estrutura,
    nome_sql,
    preencher,
)
from catalog.loader import load_catalog
from datasource.executors.fake import FakeQueryExecutor, make_result
from datasource.sql_guard import validate_sql


def _xlsx(linhas, titulo="Planilha1", formulas=()) -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = titulo
    for linha in linhas:
        aba.append(list(linha))
    for celula, formula in formulas:
        aba[celula] = formula
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


# O caso da conversa 44, reduzido: título na linha 1, linha em branco, e o
# cabeçalho na linha 3.
PAINEL = _xlsx(
    [
        ("Painel Médico - Simone (setor 3000)",),
        (),
        ("Nome do médico", "CRM", "UF de atendimento", "Potencial"),
        ("BRENDON DUTRA", "MG0104608", "BA", "A"),
        ("ELISA OLIVEIRA", "MG0049899", "MG", None),
        ("FARES NETO", "PR0023511", "MG", "B"),
    ],
    titulo="Painel",
)


@pytest.fixture(scope="module")
def catalogo():
    return load_catalog()


# ---------------------------------------------------------------- perfil


def test_cabecalho_abaixo_do_titulo():
    """A primeira linha preenchida era tomada como cabeçalho: o título
    virava a única coluna e o resto, dado."""
    estrutura = ler_estrutura("painel.xlsx", PAINEL)
    pagina = estrutura.primeira

    assert pagina.linha_do_cabecalho == 3
    assert [c.nome for c in pagina.colunas] == ["Nome do médico", "CRM", "UF de atendimento", "Potencial"]
    assert pagina.linhas == 3


def test_colunas_ganham_nome_de_sql_sem_acento():
    pagina = ler_estrutura("painel.xlsx", PAINEL).primeira

    assert [c.nome_sql for c in pagina.colunas] == ["nome_do_medico", "crm", "uf_de_atendimento", "potencial"]
    assert pagina.tabela_sql == "anexo.painel"


def test_nome_sql_de_casos_dificeis():
    assert nome_sql("  Unidades (ago/26) ") == "unidades_ago_26"
    assert nome_sql("2026") == "c_2026"
    assert nome_sql("CRM", {"crm"}) == "crm_2"
    assert nome_sql("???") == "coluna"


def test_perfil_aponta_chave_e_vazias():
    resumo = ler_estrutura("painel.xlsx", PAINEL).resumo

    assert "Tabela SQL: anexo.painel" in resumo
    assert "- B · CRM → crm · texto · 3 de 3, todas diferentes (serve de chave)" in resumo
    assert "- D · Potencial → potencial · texto · 2 de 3, 2 diferentes" in resumo


def test_registros_guardam_o_numero_da_linha_no_arquivo():
    pagina = ler_estrutura("painel.xlsx", PAINEL).primeira

    assert [numero for numero, _ in pagina.registros] == [4, 5, 6]
    assert pagina.registros[1][1][1] == "MG0049899"


def test_coluna_calculada_e_apontada():
    dados = _xlsx(
        [("Produto", "Preço", "Qtd", "Total"), ("A", 10, 2, None), ("B", 5, 3, None)],
        formulas=[("D2", "=B2*C2"), ("D3", "=B3*C3")],
    )
    pagina = ler_estrutura("vendas.xlsx", dados).primeira

    total = next(c for c in pagina.colunas if c.nome == "Total")
    assert total.formulas == 2
    assert "coluna calculada" in ler_estrutura("vendas.xlsx", dados).resumo


def test_cabecalho_com_ano_numerico_ainda_e_cabecalho():
    linhas = [("Rede", 2025, 2026), ("Pague Menos", 10, 20)]
    assert indice_do_cabecalho(linhas) == 0


def test_coluna_a_vazia_nao_vira_coluna():
    livro = Workbook()
    livro.active["B1"], livro.active["C1"] = "Rede", "Unidades"
    livro.active["B2"], livro.active["C2"] = "Pague Menos", 3
    buffer = io.BytesIO()
    livro.save(buffer)

    pagina = ler_estrutura("x.xlsx", buffer.getvalue()).primeira

    assert [c.nome for c in pagina.colunas] == ["Rede", "Unidades"]
    assert pagina.colunas[0].indice == 1


def test_csv_vira_anexo_planilha():
    dados = "Rede;Unidades\r\nPague Menos;3\r\n".encode("utf-8")
    pagina = ler_estrutura("redes.csv", dados).primeira

    assert pagina.tabela_sql == "anexo.planilha"
    assert pagina.registros == ((2, ("Pague Menos", "3")),)


# ---------------------------------------------------------------- anexo.<aba> no SQL


def test_referencias_ignoram_texto_e_comentario():
    sql = "SELECT 'anexo.falsa' -- anexo.outra\nFROM anexo.painel a JOIN anexo.\"Outra\" b ON true /* anexo.x */"

    assert anexo_sql.referencias(sql) == {"painel", "outra"}


def test_expandir_poe_os_valores_como_parametro():
    """Nenhum valor de célula entra no texto do SQL: um nome de médico com
    aspas ou ponto e vírgula é dado do unnest, não comando."""
    tabelas = anexo_sql.tabelas(ler_estrutura("painel.xlsx", PAINEL))
    sql = "SELECT a._linha, a.crm FROM anexo.painel a WHERE a.crm LIKE 'MG%'"

    expandido, params = anexo_sql.expandir(sql, tabelas)

    assert expandido.startswith("WITH anexo__painel(_linha, \"crm\") AS (SELECT * FROM unnest(%s::bigint[], %s::text[]))")
    assert "FROM anexo__painel a" in expandido
    assert "LIKE 'MG%%'" in expandido  # o % do modelo, dobrado para o driver
    assert "MG0104608" not in expandido
    assert params == [[4, 5, 6], ["MG0104608", "MG0049899", "PR0023511"]]


def test_expandir_leva_todas_as_colunas_com_estrela():
    tabelas = anexo_sql.tabelas(ler_estrutura("painel.xlsx", PAINEL))

    expandido, params = anexo_sql.expandir("SELECT * FROM (SELECT a.* FROM anexo.painel a) AS q", tabelas)

    assert len(params) == 5  # _linha + 4 colunas
    assert '"uf_de_atendimento"' in expandido


def test_sem_anexo_o_sql_passa_como_veio():
    assert anexo_sql.expandir("SELECT 1", {"painel": None}) == ("SELECT 1", None)


def test_numero_do_excel_vira_texto_sem_ponto_zero():
    """O Excel guarda o CRM 104608 como 104608.0; em coluna de texto ele
    precisa casar com '104608'."""
    assert anexo_sql._para_o_banco(104608.0, "text") == "104608"
    assert anexo_sql._para_o_banco(datetime(2026, 9, 1), "timestamp") == "2026-09-01 00:00:00"
    assert anexo_sql._para_o_banco("#N/D", "numeric") is None


def test_executor_com_anexo_registra_o_uso():
    tabelas = anexo_sql.tabelas(ler_estrutura("painel.xlsx", PAINEL))
    interno = FakeQueryExecutor([make_result(("n",), [(3,)])] * 2)
    registro = {}
    executor = anexo_sql.ExecutorComAnexo(interno, tabelas, registro)

    executor.run("SELECT 1")
    executor.run("SELECT count(*) FROM anexo.painel")

    assert interno.params[0] is None
    assert interno.params[1] == [[4, 5, 6]]
    assert registro["anexo_sql"] == ["painel"]


# ---------------------------------------------------------------- validador


def test_validador_aceita_a_aba_da_planilha(catalogo):
    guard = validate_sql("SELECT a._linha FROM anexo.painel a", catalogo, anexo={"painel"})
    assert guard.approved


def test_validador_recusa_aba_que_nao_existe(catalogo):
    guard = validate_sql("SELECT 1 FROM anexo.outra", catalogo, anexo={"painel"})
    assert not guard.approved
    assert "anexo.painel" in guard.reason


def test_validador_recusa_anexo_sem_planilha(catalogo):
    guard = validate_sql("SELECT 1 FROM anexo.painel", catalogo)
    assert not guard.approved
    assert "não tem planilha" in guard.reason


# ---------------------------------------------------------------- preenchimento pela linha


def _pedido(**campos):
    padrao = {
        "coluna_chave": "CRM",
        "chave_no_resultado": "_linha",
        "colunas": [{"coluna_destino": "Representante", "valor_no_resultado": "representante",
                     "justificativa": "território da UTC do médico"}],
        "aba": "",
    }
    padrao.update(campos)
    return PedidoDePreenchimento.do_plano(padrao)


def _linhas_da_aba(dados, aba=None):
    livro = load_workbook(io.BytesIO(dados))
    return list((livro[aba] if aba else livro.worksheets[0]).iter_rows(values_only=True))


def test_preenche_cada_linha_pelo_numero_dela():
    resultado = make_result(("_linha", "representante"), [(6, "HERMES"), (4, "ANA"), (5, None)])

    feita = preencher("painel.xlsx", PAINEL, _pedido(), resultado.columns, resultado.rows)

    linhas = _linhas_da_aba(feita.dados)
    assert linhas[2] == ("Nome do médico", "CRM", "UF de atendimento", "Potencial", "Representante")
    assert linhas[3][4] == "ANA"
    assert linhas[4][4] is None
    assert linhas[5][4] == "HERMES"
    assert (feita.preenchidas, feita.vazias_no_banco, feita.total) == (2, 1, 3)
    assert feita.criadas == ("Representante",)


def test_linha_repetida_no_resultado_fica_em_branco():
    resultado = make_result(("_linha", "representante"), [(4, "ANA"), (4, "BIA"), (5, "CAIO"), (6, "DIO")])

    feita = preencher("painel.xlsx", PAINEL, _pedido(), resultado.columns, resultado.rows)

    assert feita.ambiguas == 1
    assert _linhas_da_aba(feita.dados)[3][4] is None


def test_linha_sem_resultado_e_citada_pela_coluna_chave():
    resultado = make_result(("_linha", "representante"), [(4, "ANA")])

    feita = preencher("painel.xlsx", PAINEL, _pedido(), resultado.columns, resultado.rows)

    assert feita.sem_correspondencia_nomes == ("MG0049899", "PR0023511")


def test_coluna_existente_nao_perde_o_que_ja_tinha():
    """Preservar a base por padrão: a célula com valor fica; a vazia recebe."""
    resultado = make_result(("_linha", "potencial"), [(4, "Z"), (5, "Y"), (6, "X")])
    pedido = _pedido(colunas=[{"coluna_destino": "Potencial", "valor_no_resultado": "potencial"}])

    feita = preencher("painel.xlsx", PAINEL, pedido, resultado.columns, resultado.rows)

    linhas = _linhas_da_aba(feita.dados)
    assert [l[3] for l in linhas[3:]] == ["A", "Y", "B"]
    assert (feita.preservadas, feita.sobrescritas) == (2, 0)
    assert feita.atualizadas == ("Potencial",)


def test_sobrescrever_so_com_pedido_explicito():
    resultado = make_result(("_linha", "potencial"), [(4, "Z"), (5, "Y"), (6, "X")])
    pedido = _pedido(colunas=[{"coluna_destino": "Potencial", "valor_no_resultado": "potencial", "sobrescrever": True}])

    feita = preencher("painel.xlsx", PAINEL, pedido, resultado.columns, resultado.rows)

    assert [l[3] for l in _linhas_da_aba(feita.dados)[3:]] == ["Z", "Y", "X"]
    assert feita.sobrescritas == 2


def test_formula_da_pessoa_continua_formula():
    dados = _xlsx(
        [("Produto", "Preço", "Qtd", "Total"), ("A", 10, 2, None), ("B", 5, 3, None)],
        formulas=[("D2", "=B2*C2"), ("D3", "=B3*C3")],
    )
    resultado = make_result(("_linha", "und"), [(2, 7), (3, 8)])
    pedido = _pedido(coluna_chave="Produto", colunas=[{"coluna_destino": "Unidades", "valor_no_resultado": "und"}])

    feita = preencher("vendas.xlsx", dados, pedido, resultado.columns, resultado.rows)

    linhas = _linhas_da_aba(feita.dados)
    assert linhas[1] == ("A", 10, 2, "=B2*C2", 7)


def test_resultado_sem_linha_e_recusado():
    resultado = make_result(("crm", "representante"), [("MG0104608", "ANA")])
    with pytest.raises(AnexoRecusado):
        preencher("painel.xlsx", PAINEL, _pedido(), resultado.columns, resultado.rows)


def test_csv_preenchido_pela_linha():
    dados = "CRM;UF\r\nMG1;MG\r\nSP2;SP\r\n".encode("utf-8-sig")
    resultado = make_result(("_linha", "representante"), [(3, "BIA"), (2, "ANA")])

    feita = preencher("medicos.csv", dados, _pedido(), resultado.columns, resultado.rows)

    texto = feita.dados.decode("utf-8-sig").splitlines()
    assert texto == ["CRM;UF;Representante", "MG1;MG;ANA", "SP2;SP;BIA"]


# ---------------------------------------------------------------- abas novas e notas


def test_aba_nova_com_grafico_nao_toca_na_base():
    dados, nome, titulo, grafico = adicionar_aba(
        "painel.xlsx", PAINEL, "Resumo por UF", ("uf", "medicos"), [("MG", 2), ("BA", 1)], "barras"
    )

    livro = load_workbook(io.BytesIO(dados))
    assert livro.sheetnames == ["Painel", "Resumo por UF"]
    assert list(livro["Resumo por UF"].iter_rows(values_only=True)) == [("uf", "medicos"), ("MG", 2), ("BA", 1)]
    assert grafico and len(livro["Resumo por UF"]._charts) == 1
    assert (nome, titulo) == ("painel.xlsx", "Resumo por UF")


def test_aba_nova_com_nome_repetido_ganha_sufixo():
    _, _, titulo, _ = adicionar_aba("painel.xlsx", PAINEL, "Painel", ("a",), [(1,)])
    assert titulo == "Painel (2)"


def test_aba_nova_em_csv_vira_xlsx():
    dados = "Rede;Unidades\r\nPague Menos;3\r\n".encode("utf-8")

    novo, nome, _, _ = adicionar_aba("redes.csv", dados, "Resumo", ("total",), [(3,)])

    assert nome == "redes.xlsx"
    assert load_workbook(io.BytesIO(novo)).sheetnames == ["Dados", "Resumo"]


def test_notas_do_jarvis_substituem_as_antigas():
    notas = [{"onde": "Aba Painel · coluna Representante", "o_que": "coluna nova", "por_que": "pedido",
              "origem": "consulta", "resultado": "3 de 3"}]
    uma, nome = anotar("painel.xlsx", PAINEL, notas)
    duas, _ = anotar(nome, uma, notas)

    livro = load_workbook(io.BytesIO(duas))
    assert livro.sheetnames == ["Painel", NOME_DAS_NOTAS]
    assert livro[NOME_DAS_NOTAS]["A2"].value == "Aba Painel · coluna Representante"


# ---------------------------------------------------------------- conferência


def _permissoes(escrever=("Representante",), sobrescrever=(), novas=()):
    return {"Painel": {"escrever": set(escrever), "sobrescrever": set(sobrescrever)}, "_abas_novas": list(novas)}


def test_conferencia_aprova_o_preenchimento_e_mede_a_cobertura():
    resultado = make_result(("_linha", "representante"), [(4, "SEM REP"), (5, "SEM REP"), (6, "SEM REP")])
    feita = preencher("painel.xlsx", PAINEL, _pedido(), resultado.columns, resultado.rows)

    conferencia = qa.conferir("painel.xlsx", PAINEL, "painel.xlsx", feita.dados, _permissoes())

    assert conferencia.ok
    cobertura = conferencia.cobertura[0]
    assert (cobertura.coluna, cobertura.com_valor, cobertura.total) == ("Representante", 3, 3)
    assert (cobertura.dominante, cobertura.vezes_do_dominante) == ("SEM REP", 3)


def test_conferencia_barra_celula_original_alterada():
    livro = load_workbook(io.BytesIO(PAINEL))
    livro["Painel"]["B4"] = "ADULTERADO"
    buffer = io.BytesIO()
    livro.save(buffer)

    conferencia = qa.conferir("painel.xlsx", PAINEL, "painel.xlsx", buffer.getvalue(), _permissoes())

    assert not conferencia.ok
    assert "linha 4, coluna 2" in conferencia.problemas[0]


def test_conferencia_barra_aba_que_ninguem_pediu():
    dados, _, _, _ = adicionar_aba("painel.xlsx", PAINEL, "Intrusa", ("a",), [(1,)])

    conferencia = qa.conferir("painel.xlsx", PAINEL, "painel.xlsx", dados, _permissoes())

    assert not conferencia.ok
    assert "Intrusa" in conferencia.problemas[0]


def test_conferencia_aceita_aba_pedida_e_notas():
    dados, _, titulo, _ = adicionar_aba("painel.xlsx", PAINEL, "Resumo", ("a",), [(1,)])
    dados, nome = anotar("painel.xlsx", dados, [{"onde": "x"}])

    conferencia = qa.conferir("painel.xlsx", PAINEL, nome, dados, _permissoes(novas=[titulo]))

    assert conferencia.ok


def test_conferencia_barra_linha_que_sumiu():
    livro = load_workbook(io.BytesIO(PAINEL))
    livro["Painel"].delete_rows(6)
    buffer = io.BytesIO()
    livro.save(buffer)

    conferencia = qa.conferir("painel.xlsx", PAINEL, "painel.xlsx", buffer.getvalue(), _permissoes())

    assert not conferencia.ok
    assert "tinha 6 linhas e ficou com 5" in conferencia.problemas[0]


# ---------------------------------------------------------------- pasta de trabalho de racional


def _racional() -> bytes:
    """A forma da Metas_FV: tabela de representantes, linha de totais e,
    embaixo, o bloco de premissas que as fórmulas da tabela usam."""
    livro = Workbook()
    aba = livro.active
    aba.title = "Metas_FV "
    aba.append(["NOME CT", "SETOR", "META 1T", "META 2T"])
    aba.append(["ANA", "2713", 405, None])
    aba.append(["BIA", "-", 472, None])
    aba.append(["CAIO", "1411", 323, None])
    aba["D2"], aba["D3"], aba["D4"] = "=$B$9*C2/SUM($C$2:$C$4)", "=$B$9*C3/SUM($C$2:$C$4)", "=$B$9*C4/SUM($C$2:$C$4)"
    aba["C5"] = "=SUM(C2:C4)"
    aba["A8"], aba["B8"] = "SKU", "META 2T2026"
    aba["A9"], aba["B9"] = "ISOLADO 30", 8634
    aba["A3"].comment = __import__("openpyxl").comments.Comment("Setor vago desde agosto", "Rubens")
    custos = livro.create_sheet("AUX_Custos")
    custos.append(["Nível", "Custo"])
    custos.append([1, 0.24])
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def test_tabela_principal_para_antes_dos_totais_e_das_premissas():
    pagina = ler_estrutura("metas.xlsx", _racional()).primeira

    assert pagina.linhas == 3
    assert [n for n, _ in pagina.registros] == [2, 3, 4]
    assert ("A9", "ISOLADO 30") in pagina.fora_da_tabela and ("B9", 8634) in pagina.fora_da_tabela


def test_resumo_traz_a_formula_a_premissa_e_o_comentario():
    resumo = ler_estrutura("metas.xlsx", _racional()).resumo

    assert "- D · META 2T → meta_2t · vazia" not in resumo
    assert "calculada em 3 linhas (coluna calculada" in resumo or "=$B$9*C2/SUM($C$2:$C$4)" in resumo
    assert "L9: A=ISOLADO 30 B=8634" in resumo
    assert "A3: Setor vago desde agosto" in resumo
    assert "anexo.metas_fv_celulas" in resumo


def test_celulas_viram_tabela_consultavel():
    tabelas = anexo_sql.tabelas(ler_estrutura("metas.xlsx", _racional()))

    assert {"metas_fv", "metas_fv_celulas", "aux_custos", "aux_custos_celulas"} <= set(tabelas)
    expandido, params = anexo_sql.expandir(
        "SELECT c.numero, c.formula FROM anexo.metas_fv_celulas c WHERE c._linha = 9 AND c.coluna = 'B'", tabelas
    )
    assert "anexo__metas_fv_celulas" in expandido
    linhas, colunas, numeros = params[0], params[1], params[2]
    assert (9, "B", 8634) in zip(linhas, colunas, numeros)
    formulas = dict(((l, c), f) for l, c, f in zip(params[0], params[1], params[3]))
    assert formulas[(2, "D")] == "=$B$9*C2/SUM($C$2:$C$4)"


def test_teto_do_resumo_cresce_com_as_abas():
    from attachments.limites import MAX_CARACTERES_DO_RESUMO

    estrutura = ler_estrutura("metas.xlsx", _racional())

    assert estrutura.teto > MAX_CARACTERES_DO_RESUMO
    assert "(resumo cortado no limite)" not in estrutura.resumo


def test_exemplo_longo_e_encurtado():
    texto = "visita por telefone, " * 50
    dados = _xlsx([("Nome", "Comentário", "UF"), ("ANA", texto, "MG"), ("BIA", texto, "SP")])

    resumo = ler_estrutura("x.xlsx", dados).resumo

    assert "- C · UF → uf" in resumo
    assert len(resumo) < 1500


def test_preenchimento_nao_conta_o_bloco_de_premissas():
    resultado = make_result(("_linha", "gr"), [(2, "IVAN"), (3, "GABRIEL"), (4, "IVAN")])
    pedido = _pedido(coluna_chave="NOME CT", colunas=[{"coluna_destino": "GR", "valor_no_resultado": "gr"}], aba="Metas_FV ")

    feita = preencher("metas.xlsx", _racional(), pedido, resultado.columns, resultado.rows)

    assert (feita.preenchidas, feita.total, feita.sem_correspondencia) == (3, 3, 0)


def test_aba_com_o_nome_do_arquivo_cai_na_unica_aba():
    """O modelo pôs no campo `aba` o nome do ARQUIVO, e nada foi preenchido
    (Painel, 2026-10-01). Com uma aba só, não há o que errar."""
    resultado = make_result(("_linha", "representante"), [(4, "ANA")])

    feita = preencher("painel.xlsx", PAINEL, _pedido(aba="Painel Médico - Simone (setor 3000)"),
                      resultado.columns, resultado.rows)

    assert feita.preenchidas == 1


def test_conferencia_tolera_o_17o_algarismo_e_pega_mudanca_de_verdade():
    """O openpyxl regrava 167.26185324474665 como 167.2618532447466: não é
    mudança (o Excel guarda 15 algarismos), e barrou a aba nova da planilha
    de metas. Mudança de verdade continua barrada."""
    assert qa._mesmo(167.26185324474665, 167.2618532447466)
    assert qa._mesmo(0.020732466089820132, 0.02073246608982013)
    assert not qa._mesmo(167.2618, 167.2619)
    assert not qa._mesmo(0, 0.0000001)


def test_conferencia_aprova_a_planilha_de_metas_regravada():
    """A planilha real de metas, aberta e salva pelo openpyxl com uma aba a
    mais: nenhuma célula original pode ser acusada."""
    from pathlib import Path

    arquivo = Path(__file__).resolve().parents[2] / "jarvis_analise_xslx_csv" / "Racional Metas 2T2026 (V2.3).xlsx"
    if not arquivo.exists():
        pytest.skip("planilha de exemplo fora do repositório")
    original = arquivo.read_bytes()
    dados, nome, titulo, _ = adicionar_aba(arquivo.name, original, "Resumo", ("a",), [(1,)])

    conferencia = qa.conferir(arquivo.name, original, nome, dados, {"_abas_novas": [titulo]})

    assert conferencia.ok, conferencia.problemas


def test_update_so_das_vazias_nao_conta_as_ja_preenchidas_como_faltando():
    """A consulta pediu só as linhas com Potencial vazio; as preenchidas
    ficaram de fora de propósito e não são "sem correspondência"."""
    resultado = make_result(("_linha", "potencial"), [(5, "M")])
    pedido = _pedido(colunas=[{"coluna_destino": "Potencial", "valor_no_resultado": "potencial"}])

    feita = preencher("painel.xlsx", PAINEL, pedido, resultado.columns, resultado.rows)

    assert (feita.preenchidas, feita.total, feita.sem_correspondencia) == (1, 1, 0)
