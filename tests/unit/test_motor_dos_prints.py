"""Motor 3.0 dos prints (ADR-0035): o print como dado.

Os casos são os prints que chegaram ao Jarvis em produção até 2026-10-08: a
lista de 230 CRMs que virou uma tira ilegível (conversa 44), a faixa de
planilha "ilegível" que mesmo assim virou pergunta sobre uma planilha que não
existia (75), a tabela a completar (7 e 8), o "do que se trata esse print?"
(80) — e o que o motor passa a fazer com eles.
"""

import io

import pytest
from openpyxl import load_workbook
from PIL import Image

from ai_orchestrator import canned
from ai_orchestrator.models import AIReply
from ai_orchestrator.orchestrator import handle_message
from ai_orchestrator.providers.base import BlocoLido, ImageRequest
from ai_orchestrator.providers.openai_provider import _pedaco_do_print
from attachments import deposito, transcricao
from attachments.imagem import pedacos, preparar
from catalog.loader import load_catalog
from conversations.models import Conversation
from datasource.executors.fake import FakeQueryExecutor, make_result
from messaging.channels.fake import FakeChannel
from messaging.models import Message
from tests.fakes.providers import ScriptedAIProvider, leitura, plano, resposta

pytestmark = pytest.mark.django_db

CRMS = [f"MG{n:07d}" for n in range(104600, 104630)]


@pytest.fixture(scope="module")
def catalogo():
    return load_catalog()


@pytest.fixture
def conversa(django_user_model):
    user = django_user_model.objects.create_user("paulo", password="x")
    return Conversation.objects.create(user=user)


def _png(largura, altura) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (largura, altura), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _print(conversa, texto, dados=None, client_id="c-1"):
    preparada = preparar(dados or _png(1200, 700))
    return Message.objects.create(
        conversation=conversa, direction=Message.Direction.INBOUND, content=texto,
        client_message_id=client_id, status=Message.Status.RECEIVED,
        anexo_tipo=Message.Anexo.IMAGEM, anexo_nome="print-colado.png",
        anexo_resumo="imagem", anexo_token=deposito.guardar(preparada.dados),
    )


def _pergunta(conversa, texto, client_id):
    return Message.objects.create(
        conversation=conversa, direction=Message.Direction.INBOUND, content=texto,
        client_message_id=client_id, status=Message.Status.RECEIVED,
    )


def _responder(mensagem, catalogo, provider, executor=None):
    return handle_message(mensagem, channel=FakeChannel(), provider=provider,
                          executor=executor or FakeQueryExecutor([]), catalog=catalogo)


def _lista(crms, titulo=""):
    return BlocoLido(tipo="lista", titulo=titulo, colunas=("CRM",), linhas=tuple((c,) for c in crms))


VENDAS = BlocoLido(tipo="tabela", colunas=("Rede", "Unidades"), linhas=(
    ("Pague Menos", "1.234"), ("Drogasil", "766"), ("Total", "2.000"),
))


# ------------------------------------------------ leitura em pedaços


def test_lista_alta_e_lida_em_pedacos_legiveis_e_nao_encolhida():
    """Conversa 44: 230 CRMs num print alto viraram uma tira de 85 px de
    largura ao reduzir para 1.280 no lado maior — "resolução baixa demais"."""
    preparada = preparar(_png(400, 6000))
    partes, direcao = pedacos(preparada.dados)

    assert (preparada.largura, preparada.altura) == (400, 6000)       # nada encolheu
    assert direcao == "vertical" and len(partes) == preparada.pedacos == 3
    assert all(Image.open(io.BytesIO(p)).size[0] == 400 for p in partes)


def test_faixa_larga_nao_e_reduzida_a_ponto_de_apagar_o_texto():
    """Conversa 75: a faixa larga da planilha chegou "ilegível"."""
    preparada = preparar(_png(3000, 300))

    assert preparada.altura == 300
    assert pedacos(preparada.dados)[1] == "horizontal"


def test_print_comum_continua_numa_leitura_so():
    preparada = preparar(_png(1920, 1080))

    assert preparada.pedacos == 1
    assert pedacos(preparada.dados) == ([preparada.dados], "")


def test_pedaco_seguinte_recebe_os_cabecalhos_do_primeiro():
    texto = _pedaco_do_print(ImageRequest(question="q", imagem_png=b"", pedaco=2, pedacos=3,
                                          direcao="vertical", colunas_lidas=("CRM", "Nome")))

    assert "Pedaço 2 de 3" in texto and "CRM | Nome" in texto
    assert _pedaco_do_print(ImageRequest(question="q", imagem_png=b"")) == ""


# ------------------------------------------------ junção e conferência


def test_juntar_pedacos_tira_a_linha_repetida_na_emenda():
    primeiro = [_lista(CRMS[:12])]
    segundo = [_lista(CRMS[10:22])]       # duas linhas na sobreposição
    terceiro = [BlocoLido(tipo="lista", colunas=("MG0104622",), linhas=tuple((c,) for c in CRMS[23:]))]

    blocos, avisos = transcricao.juntar([primeiro, segundo, terceiro], "vertical")

    # O terceiro pedaço não viu cabeçalho: a primeira linha dele é dado.
    assert [l[0] for l in blocos[0]["linhas"]] == CRMS and avisos == []


def test_juntar_faixa_larga_poe_as_colunas_lado_a_lado():
    esquerda = [BlocoLido(tipo="tabela", colunas=("UF", "Cidade"), linhas=(("MG", "BH"), ("SP", "Santos")))]
    direita = [BlocoLido(tipo="tabela", colunas=("Cidade", "GR"), linhas=(("BH", "Ivan"), ("Santos", "Gabriel")))]

    blocos, avisos = transcricao.juntar([esquerda, direita], "horizontal")

    assert blocos[0]["colunas"] == ["UF", "Cidade", "GR"]
    assert blocos[0]["linhas"][1] == ["SP", "Santos", "Gabriel"] and not avisos


def test_faixa_que_nao_alinha_fica_com_a_parte_da_esquerda_e_avisa():
    esquerda = [BlocoLido(tipo="tabela", colunas=("UF",), linhas=(("MG",), ("SP",)))]
    direita = [BlocoLido(tipo="tabela", colunas=("GR",), linhas=(("Ivan",),))]

    blocos, avisos = transcricao.juntar([esquerda, direita], "horizontal")

    assert blocos[0]["colunas"] == ["UF"] and "não alinharam" in avisos[0]


def test_conferencia_tipa_os_numeros_e_tira_o_total_que_fecha():
    t = transcricao.conferir([VENDAS])

    [bloco] = t.blocos
    assert bloco.linhas == (("Pague Menos", 1234.0), ("Drogasil", 766.0))
    assert "fecha com a soma" in t.avisos[0]


def test_total_que_nao_fecha_e_dito():
    errado = BlocoLido(tipo="tabela", colunas=("Rede", "Unidades"),
                       linhas=(("A", "10"), ("B", "20"), ("Total", "40")))

    assert "não fecha" in transcricao.conferir([errado]).avisos[0]


def test_celula_ilegivel_fica_em_branco_e_e_citada_pelo_nome():
    bloco = BlocoLido(tipo="tabela", colunas=("Rede", "Unidades"), linhas=(("Pague Menos", "?"), ("Drogasil", "7")))

    t = transcricao.conferir([bloco])

    assert t.blocos[0].linhas[0] == ("Pague Menos", None)
    assert "Pague Menos" in t.avisos[0]


@pytest.mark.parametrize("texto,valor", [
    ("1.234,56", 1234.56), ("R$ 1.234", 1234.0), ("12,5%", 12.5), ("-3,2", -3.2),
    ("(120)", -120.0), ("2.461", 2461.0), ("1.5", 1.5), ("3,2 mil", 3200.0), ("Drogasil", None),
])
def test_numero_escrito_como_o_brasil_escreve(texto, valor):
    assert transcricao.numero_br(texto) == valor


def test_codigo_repetido_na_lista_e_avisado():
    """Teste real de 2026-10-09: 227 CRMs lidos, 2 repetidos."""
    lista = _lista(["MG0104608", "MG0049899", "MG0104608"])

    t = transcricao.conferir([lista])

    assert len(t.blocos[0].linhas) == 3
    assert "MG0104608" in t.avisos[0] and "mais de uma vez" in t.avisos[0]


def test_ultimo_pedaco_nao_recua_e_repete_linhas():
    """Teste real de 2026-10-09: o último pedaço recuava até caber inteiro e
    repetia ~65 linhas do anterior (224 CRMs viraram 290)."""
    preparada = preparar(_png(131, 4726))
    partes, _ = pedacos(preparada.dados)

    alturas = [Image.open(io.BytesIO(p)).size[1] for p in partes]
    assert alturas == [2048, 2048, 726]


def test_pedido_longo_do_leitor_nao_entra_na_pergunta(conversa, catalogo):
    """Teste real de 2026-10-09: o leitor listou 96 CRMs em `pergunta_ao_banco`
    e o planejador os copiou num VALUES, ignorando os 227 de `anexo.<aba>`."""
    provider = ScriptedAIProvider(
        [plano(sql="SELECT 1")], [resposta("Feito.")],
        leituras=[leitura("Lista.", blocos=(_lista(CRMS),), operacao="cruzar", precisa_do_banco=True,
                          pergunta_ao_banco="Traga o representante dos CRMs " + ", ".join(CRMS * 3))],
    )

    _responder(_print(conversa, "Traga o representante de cada CRM"), catalogo, provider,
               FakeQueryExecutor([make_result(("x",), [(1,)])]))

    assert provider.plan_requests[0].question == "Traga o representante de cada CRM"


def test_crm_e_codigo_continuam_texto():
    lista = BlocoLido(tipo="lista", colunas=("CRM",), linhas=(("MG0104608",), ("0012345",)))

    assert transcricao.conferir([lista]).blocos[0].linhas[0] == ("MG0104608",)


# ------------------------------------------------ rotas do orquestrador


def test_lista_de_crms_do_print_vira_tabela_e_cruza_com_o_banco(conversa, catalogo):
    """Conversa 44: com a lista transcrita, o pedido vira uma consulta sobre
    `anexo.<aba>` — sem a pessoa colar os CRMs no texto."""
    provider = ScriptedAIProvider(
        [plano(sql="SELECT 1")], [resposta("Cruzei os CRMs do print.")],
        leituras=[
            leitura("Lista de CRMs.", blocos=(_lista(CRMS[:16]),), operacao="cruzar", precisa_do_banco=True),
            leitura("continuação", blocos=(_lista(CRMS[14:]),)),
            leitura("continuação", blocos=()),
        ],
    )
    mensagem = _print(conversa, "Traga o representante de cada CRM desta lista", _png(400, 6000))

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([make_result(("x",), [(1,)])]))

    assert len(provider.image_requests) == 3
    assert not provider.plan_requests[0].so_a_planilha      # cruzar precisa das regras do banco
    assert provider.image_requests[1].colunas_lidas == ("CRM",)
    planilha = provider.plan_requests[0].planilha
    assert planilha.startswith("ESTA PLANILHA É A TRANSCRIÇÃO DE UM PRINT")
    assert "30 linhas" in planilha and "anexo." in planilha
    reply = AIReply.objects.get(message=mensagem)
    assert reply.raw_response["imagem"]["virou"] == "planilha"
    assert reply.raw_response["print_transcrito"]["token"]


def test_print_ilegivel_pede_de_outro_jeito_sem_inventar_pergunta(conversa, catalogo):
    """Conversa 75: "ilegível", e mesmo assim saiu uma pergunta sobre "a
    planilha enviada" — US$ 0,19 para pedir os nomes das colunas de uma
    planilha que não existia."""
    provider = ScriptedAIProvider(leituras=[leitura(
        "Faixa de planilha em resolução baixa.", legivel=False, motivo_ilegivel="a faixa está pequena demais",
        precisa_do_banco=True, pergunta_ao_banco="Vincular cada estado e cidade da planilha enviada ao GR",
    )])

    reply = _responder(_print(conversa, "vincular os dados colados ao GR"), catalogo, provider)

    assert reply.rule == "print_ilegivel"
    assert reply.reply_text.startswith("Não consegui ler o print com segurança: a faixa está pequena demais.")
    assert provider.plan_requests == []


def test_conta_sobre_o_print_vai_para_o_sql(conversa, catalogo):
    """"Qual o total?" é conta: o leitor não a faz, e a resposta vem da
    consulta sobre a transcrição, com a ancoragem de sempre."""
    provider = ScriptedAIProvider(
        [plano(sql="SELECT 1")], [resposta("Foram 2000 unidades.")],
        leituras=[leitura("Tabela de vendas.", blocos=(VENDAS,), operacao="analisar")],
    )

    _responder(_print(conversa, "qual o total de unidades?"), catalogo, provider,
               FakeQueryExecutor([make_result(("total",), [(2000.0,)])]))

    assert len(provider.plan_requests) == 1
    # Conta só sobre o print: sem o documento de negócio (~US$ 0,09 → ~0,02).
    assert provider.plan_requests[0].so_a_planilha


def test_descricao_com_numero_que_nao_esta_no_print_vai_para_o_sql(conversa, catalogo):
    """O leitor somou por conta própria (1.234 + 766): número que não está na
    transcrição não sai sem conferência."""
    provider = ScriptedAIProvider(
        [plano(sql="SELECT 1")], [resposta("Somam 2000 unidades.")],
        leituras=[leitura("Somam 2.000 unidades, 61,7% na Pague Menos.", blocos=(VENDAS,), operacao="descrever")],
    )

    _responder(_print(conversa, "o que esse print mostra?"), catalogo, provider,
               FakeQueryExecutor([make_result(("total",), [(2000.0,)])]))

    assert len(provider.plan_requests) == 1


def test_descricao_ancorada_fica_no_leitor_barato_e_o_print_segue_na_conversa(conversa, catalogo):
    provider = ScriptedAIProvider(
        [plano(sql="SELECT 1")], [resposta("Feito.")],
        leituras=[leitura("A Pague Menos vendeu 1.234 unidades e a Drogasil 766.",
                          blocos=(VENDAS,), operacao="descrever")],
    )
    reply = _responder(_print(conversa, "o que esse print mostra?"), catalogo, provider)

    assert reply.decision == AIReply.Decision.IMAGE_READING
    assert "fecha com a soma" in reply.reply_text
    assert provider.plan_requests == []

    # "e quanto isso é do total de setembro?": o print continua como tabela.
    _responder(_pergunta(conversa, "e quanto a Pague Menos vendeu em setembro?", "c-2"), catalogo, provider,
               FakeQueryExecutor([make_result(("x",), [(1,)])]))
    assert "TRANSCRIÇÃO DE UM PRINT" in provider.plan_requests[0].planilha


def test_print_vira_excel_sem_planejador(conversa, catalogo):
    provider = ScriptedAIProvider(leituras=[leitura("Tabela de vendas.", blocos=(VENDAS,), operacao="transformar")])
    mensagem = _print(conversa, "passa esse print para Excel")

    reply = _responder(mensagem, catalogo, provider)

    assert reply.rule == "print_em_planilha" and provider.plan_requests == []
    mensagem.refresh_from_db()
    assert mensagem.anexo_resposta_nome == "print-colado.xlsx"
    livro = load_workbook(io.BytesIO(deposito.buscar(mensagem.anexo_resposta_token)))
    assert livro.sheetnames == ["Tabela do print", "Notas do Jarvis"]
    assert list(livro["Tabela do print"].iter_rows(values_only=True))[1] == ("Pague Menos", 1234)
    assert "Baixar planilha preenchida" in reply.reply_text


def test_print_sem_dado_continua_so_lido(conversa, catalogo):
    """Conversa 80: "Do que se trata esse print?" com a tela do Jarvis."""
    provider = ScriptedAIProvider(leituras=[leitura("A tela inicial do Jarvis.", tipo_do_print="tela")])

    reply = _responder(_print(conversa, "Do que se trata esse print?"), catalogo, provider)

    assert reply.decision == AIReply.Decision.IMAGE_READING
    assert provider.plan_requests == []
    assert "print_transcrito" not in reply.raw_response


def test_canned_do_print_ilegivel_diz_o_que_fazer():
    assert "arquivo original" in canned.PRINT_ILEGIVEL
