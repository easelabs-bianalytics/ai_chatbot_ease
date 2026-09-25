"""Planilha no WhatsApp e no chat web (conversa 32, 2026-09-24).

A pessoa mandou `dados_preencher.xlsx` (duas abas, cabeçalhos claros) com
"O que existe nessa planilha?". O resumo foi lido certo, mas:

1. a descrição foi trocada pelo texto de reserva ("Não consegui confirmar
   esse número…"), porque a trava de número sem fonte não via o resumo;
2. na mensagem seguinte ("Mas você conseguiu ver o que existe dentro
   dela?") a planilha já não estava — o Jarvis disse que não havia nenhuma;
3. "Preencha as duas" não tinha caminho: só uma aba por vez.
"""

import io

import pytest
from openpyxl import Workbook, load_workbook

from ai_orchestrator.models import AIReply
from attachments import deposito
from datasource.executors.fake import FakeQueryExecutor, make_result
from messaging.models import Message
from tests.fakes.providers import ScriptedAIProvider, plano, resposta
from tests.unit.test_orchestrator_anexos import _pergunta_com_anexo, _responder, catalogo, conversa  # noqa: F401

pytestmark = pytest.mark.django_db


def _duas_abas() -> bytes:
    """A planilha da conversa 32, como ela é."""
    livro = Workbook()
    redes = livro.active
    redes.title = "Planilha3"
    for linha in (("Rede", "UNIDADES SELL OUT AGO/26"), ("RAIA DROGASIL", None), ("PANVEL", None),
                  ("PAGUE MENOS", None), ("INDIANA", None)):
        redes.append(linha)
    reps = livro.create_sheet("Planilha1")
    for linha in (("REPRESENTANTE", "PX EASE YTD 2026", "PX ACHE YTD 2026"), ("MARTA ELOISA", None, None),
                  ("JONATHAN ABRAHAO", None, None), ("ALEXANDRE CIMINI", None, None)):
        reps.append(linha)
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


PLANILHA = _duas_abas()
REDES = make_result(("rede", "unidades"), [("RAIA DROGASIL", 3510.0), ("PANVEL", 812.0),
                                           ("PAGUE MENOS", 1204.0), ("INDIANA", 97.0)])
REPS = make_result(("representante", "px_ease", "px_ache"), [("MARTA ELOISA", 410.0, 380.0),
                                                            ("JONATHAN ABRAHAO", 520.0, 610.0),
                                                            ("ALEXANDRE CIMINI", 300.0, 290.0)])
SQL_REDES = "SELECT rede, SUM(und) AS unidades FROM cddd.vendas_consolidado GROUP BY 1"
SQL_REPS = "SELECT canal AS representante, SUM(und) AS px_ease, SUM(und) AS px_ache FROM cddd.vendas_consolidado GROUP BY 1"


def _mensagem(conversa, texto):
    return _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "dados_preencher.xlsx", texto=texto)


def test_descrever_a_planilha_nao_vira_texto_de_reserva(conversa, catalogo):
    """Os números da descrição ("4 redes", "AGO/26") vêm do resumo da
    planilha: são fonte, e a resposta vale."""
    descricao = ("A planilha tem 2 abas. A Planilha3 lista 4 redes (RAIA DROGASIL, PANVEL, PAGUE MENOS e "
                 "INDIANA) com UNIDADES SELL OUT AGO/26 vazia. A Planilha1 traz 3 representantes com o PX "
                 "EASE YTD 2026 e o PX ACHE YTD 2026 a preencher. Quer que eu preencha as duas?")
    provider = ScriptedAIProvider([plano("", intent="conversation", user_message=descricao)])

    reply = _responder(_mensagem(conversa, "O que existe nessa planilha?"), catalogo, provider)

    assert reply.rule != "conversa_com_numero_sem_fonte"
    assert reply.reply_text == descricao


def test_a_planilha_acompanha_a_pergunta_seguinte(conversa, catalogo):
    """"Mas você conseguiu ver o que existe dentro dela?" chega sem anexo:
    a planilha da mensagem anterior segue, avisada como tal."""
    _responder(_mensagem(conversa, "O que existe nessa planilha?"), catalogo,
               ScriptedAIProvider([plano("", intent="conversation", user_message="Tem duas abas.")]))
    seguinte = Message.objects.create(conversation=conversa, direction=Message.Direction.INBOUND,
                                      content="Mas você conseguiu ver o que existe dentro dela?", client_message_id="c-2")
    provider = ScriptedAIProvider([plano("", intent="conversation", user_message="Sim: duas abas.")])

    _responder(seguinte, catalogo, provider)

    enviado = provider.plan_requests[0].planilha
    assert "Planilha3" in enviado and "Planilha1" in enviado
    assert "mensagem anterior desta conversa" in enviado


def test_planilha_para_de_seguir_quando_a_conversa_muda_de_assunto(conversa, catalogo):
    _responder(_mensagem(conversa, "O que existe nessa planilha?"), catalogo,
               ScriptedAIProvider([plano("", intent="conversation", user_message="Tem duas abas.")]))
    outra = Message.objects.create(conversation=conversa, direction=Message.Direction.INBOUND,
                                   content="quanto vendemos ontem?", client_message_id="c-2")
    reply = _responder(outra, catalogo, ScriptedAIProvider([plano()], [resposta("A Raia Drogasil vendeu 3510 unidades.")]),
                       FakeQueryExecutor([REDES]))
    # A pergunta era outra: a resposta não ganha aviso de planilha.
    assert "planilha" not in reply.reply_text.lower()
    terceira = Message.objects.create(conversation=conversa, direction=Message.Direction.INBOUND,
                                      content="e na semana?", client_message_id="c-3")
    provider = ScriptedAIProvider([plano()], [resposta("A Raia Drogasil vendeu 3510 unidades.")])

    _responder(terceira, catalogo, provider, FakeQueryExecutor([REDES]))

    assert provider.plan_requests[0].planilha == ""


def test_preencha_as_duas_preenche_cada_aba_com_a_sua_consulta(conversa, catalogo):
    """Uma consulta por aba — chaves diferentes, rede numa e representante
    na outra — e um arquivo só com as duas preenchidas."""
    consultas = (
        {"titulo": "Sell out por rede", "sql": SQL_REDES, "reference_query_id": "B01", "preenchimento": {
            "aba": "Planilha3", "coluna_chave": "Rede", "chave_no_resultado": "rede",
            "colunas": [{"coluna_destino": "UNIDADES SELL OUT AGO/26", "valor_no_resultado": "unidades"}]}},
        {"titulo": "PX por representante", "sql": SQL_REPS, "reference_query_id": "A21", "preenchimento": {
            "aba": "Planilha1", "coluna_chave": "REPRESENTANTE", "chave_no_resultado": "representante",
            "colunas": [{"coluna_destino": "PX EASE YTD 2026", "valor_no_resultado": "px_ease"},
                        {"coluna_destino": "PX ACHE YTD 2026", "valor_no_resultado": "px_ache"}]}},
    )
    provider = ScriptedAIProvider(
        [plano("", consultas=consultas, entendimento="Preencher as duas abas")],
        [resposta("Raia Drogasil lidera com 3510 unidades.")],
    )
    mensagem = _mensagem(conversa, "Preencha as duas, por favor")

    # entregas (redes, reps) e o preenchimento com o limite alto (redes, reps)
    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([REDES, REPS, REDES, REPS]))

    assert reply.decision == AIReply.Decision.ANSWERED
    mensagem.refresh_from_db()
    livro = load_workbook(io.BytesIO(deposito.buscar(mensagem.anexo_resposta_token)))
    redes = list(livro["Planilha3"].iter_rows(values_only=True))
    reps = list(livro["Planilha1"].iter_rows(values_only=True))
    assert redes[1] == ("RAIA DROGASIL", 3510.0) and redes[4] == ("INDIANA", 97.0)
    assert reps[3] == ("ALEXANDRE CIMINI", 300.0, 290.0)
    assert "da aba `Planilha3`" in reply.reply_text and "da aba `Planilha1`" in reply.reply_text
    assert "Baixar planilha preenchida" in reply.reply_text


def test_provider_le_o_preenchimento_de_cada_consulta_so_com_planilha():
    from ai_orchestrator.providers.openai_provider import PlanoEstruturado, _consultas

    bruto = PlanoEstruturado.model_validate({
        "intent": "answer_with_data",
        "consultas": [{"titulo": "Redes", "sql": SQL_REDES, "preenchimento": {
            "aba": "Planilha3", "coluna_chave": "Rede", "chave_no_resultado": "rede",
            "colunas": [{"coluna_destino": "UNIDADES SELL OUT AGO/26", "valor_no_resultado": "unidades"}]}}],
    })

    com_planilha = _consultas(bruto, planilha=True)
    assert com_planilha[0]["preenchimento"]["aba"] == "Planilha3"
    # Sem planilha anexada, preencher não existe como caminho.
    assert "preenchimento" not in _consultas(bruto, planilha=False)[0]
