"""Print que precisa do banco, e a planilha em toda chamada ao planejador (ADR-0024).

Achados de 2026-09-21:

- o print de uma tabela vazia com "preencha para mim" era só lido, e a
  resposta dizia que não havia números na imagem;
- a chamada com o documento inteiro saía sem a planilha, e o modelo
  respondia "não há planilha anexada".
"""

import io

import pytest
from openpyxl import load_workbook

from ai_orchestrator.context import PEDIDO_DE_SECAO
from ai_orchestrator.models import AIReply
from attachments import deposito
from datasource.executors.fake import FakeQueryExecutor
from messaging.models import Message
from tests.fakes.providers import ScriptedAIProvider, leitura, plano, resposta
from tests.unit.test_orchestrator_anexos import (  # noqa: F401
    PLANILHA,
    PREENCHIMENTO,
    RESULTADO,
    _pergunta_com_anexo,
    _responder,
    catalogo,
    conversa,
)

pytestmark = pytest.mark.django_db

LINHAS_DO_PRINT = (("Pague Menos", ""), ("Drogasil", ""), ("Sem Venda", ""))


def test_print_de_tabela_vira_planilha_e_e_preenchido_com_o_banco(conversa, catalogo):
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.IMAGEM, b"png", "print.png", "preencha para mim")
    provider = ScriptedAIProvider(
        [plano(preenchimento=PREENCHIMENTO)],
        [resposta("foram 47 unidades")],
        leituras=[leitura(precisa_do_banco=True, tabela_colunas=("Rede", "Unidades"), tabela_linhas=LINHAS_DO_PRINT)],
    )

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO, RESULTADO]))

    assert reply.decision == AIReply.Decision.ANSWERED
    # O planejador recebeu a tabela do print como planilha — a forma, não números.
    assert "Rede" in provider.plan_requests[0].planilha
    mensagem.refresh_from_db()
    preenchida = deposito.buscar(mensagem.anexo_resposta_token)
    linhas = list(load_workbook(io.BytesIO(preenchida)).active.iter_rows(values_only=True))
    assert linhas[1] == ("Pague Menos", 47.0)
    assert linhas[2] == ("Drogasil", 30.0)
    assert reply.raw_response["imagem"]["virou"] == "planilha"
    # A resposta veio do banco: a decisão não é a de leitura de imagem.
    assert reply.decision != AIReply.Decision.IMAGE_READING


def test_print_que_pede_conferencia_vira_pergunta_ao_banco(conversa, catalogo):
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.IMAGEM, b"png", "print.png", "isso bate?")
    pergunta = "Quantas unidades a Pague Menos vendeu em agosto de 2026?"
    provider = ScriptedAIProvider(
        [plano()], [resposta("foram 47 unidades")],
        leituras=[leitura(precisa_do_banco=True, pergunta_ao_banco=pergunta)],
    )

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    assert provider.plan_requests[0].question == pergunta
    assert provider.plan_requests[0].planilha == ""
    assert reply.decision == AIReply.Decision.ANSWERED
    assert reply.raw_response["imagem"]["pergunta"] == pergunta


def test_imagem_convertida_nao_altera_a_pergunta_gravada(conversa, catalogo):
    """A conversão é só em memória: no banco fica o que a pessoa escreveu,
    com a imagem."""
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.IMAGEM, b"png", "print.png", "isso bate?")
    provider = ScriptedAIProvider(
        [plano()], [resposta("foram 47 unidades")],
        leituras=[leitura(precisa_do_banco=True, pergunta_ao_banco="outra pergunta")],
    )

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    gravada = Message.objects.get(pk=mensagem.pk)
    assert (gravada.content, gravada.anexo_tipo) == ("isso bate?", Message.Anexo.IMAGEM)


def test_precisa_do_banco_sem_tabela_nem_pergunta_fica_na_leitura(conversa, catalogo):
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.IMAGEM, b"png", "print.png")
    provider = ScriptedAIProvider(leituras=[leitura("Não identifiquei o que buscar.", precisa_do_banco=True)])

    reply = _responder(mensagem, catalogo, provider)

    assert reply.decision == AIReply.Decision.IMAGE_READING
    assert provider.plan_requests == []


def test_planilha_acompanha_a_chamada_do_documento_inteiro(conversa, catalogo):
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    pede_secao = plano(sql="", intent="unknown", reason=f"{PEDIDO_DE_SECAO}: sell_out")
    provider = ScriptedAIProvider(
        [pede_secao, plano(preenchimento=PREENCHIMENTO)], [resposta("foram 47 unidades")]
    )

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO, RESULTADO]))

    assert provider.plan_requests[1].full_context is True
    assert provider.plan_requests[1].planilha == mensagem.anexo_resumo


def test_planilha_acompanha_a_correcao_da_consulta(conversa, catalogo):
    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    provider = ScriptedAIProvider(
        [plano(sql="DELETE FROM cddd.pdvs"), plano(preenchimento=PREENCHIMENTO)],
        [resposta("foram 47 unidades")],
    )

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO, RESULTADO]))

    assert provider.plan_requests[1].error_note
    assert provider.plan_requests[1].planilha == mensagem.anexo_resumo


# ------------------------------------------------------- planilha pendente
# Teste real de 2026-09-21: com a planilha, o Jarvis perguntou "painel atual
# ou território?". A resposta da pessoa chegava sem anexo, e o preenchimento
# nunca acontecia.


def _seguimento(conversa, texto="painel atual", client_id="c-2"):
    return Message.objects.create(
        conversation=conversa, direction=Message.Direction.INBOUND, content=texto, client_message_id=client_id,
    )


def _pede_detalhe():
    return plano(sql="", intent="clarify", clarification_question="Painel atual ou território?")


def test_planilha_espera_a_resposta_ao_pedido_de_detalhe(conversa, catalogo):
    primeira = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    _responder(primeira, catalogo, ScriptedAIProvider([_pede_detalhe()]))
    segunda = _seguimento(conversa)
    provider = ScriptedAIProvider([plano(preenchimento=PREENCHIMENTO)], [resposta("foram 47 unidades")])

    reply = _responder(segunda, catalogo, provider, FakeQueryExecutor([RESULTADO, RESULTADO]))

    assert provider.plan_requests[0].planilha == primeira.anexo_resumo
    segunda.refresh_from_db()
    assert deposito.buscar(segunda.anexo_resposta_token) is not None
    assert reply.raw_response["planilha_herdada_de"] == primeira.pk
    # No banco a resposta da pessoa continua sem anexo.
    assert segunda.anexo_tipo == ""


def test_pedido_de_detalhe_renova_o_prazo_da_planilha(conversa, catalogo, monkeypatch):
    prazos = []
    monkeypatch.setattr(deposito, "prolongar", lambda token, segundos: prazos.append(segundos) or True)
    primeira = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")

    reply = _responder(primeira, catalogo, ScriptedAIProvider([_pede_detalhe()]))

    assert prazos == [30 * 60]
    assert reply.raw_response["planilha_pendente"]["token"] == primeira.anexo_token


def test_planilha_nao_reaparece_se_a_conversa_mudou_de_assunto(conversa, catalogo):
    primeira = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    _responder(primeira, catalogo, ScriptedAIProvider([_pede_detalhe()]))
    outra = _seguimento(conversa, "quanto vendemos ontem?", "c-2")
    _responder(outra, catalogo, ScriptedAIProvider([plano()], [resposta("foram 47 unidades")]),
               FakeQueryExecutor([RESULTADO]))
    terceira = _seguimento(conversa, "e na semana?", "c-3")
    provider = ScriptedAIProvider([plano()], [resposta("foram 47 unidades")])

    _responder(terceira, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    assert provider.plan_requests[0].planilha == ""


def test_correcao_usa_o_mesmo_contexto_do_plano_corrigido(conversa, catalogo):
    """Plano feito com o documento inteiro e recusado pelo validador voltava
    ao contexto recortado na correção, e o modelo desistia (2026-09-21)."""
    from ai_orchestrator.providers.base import AIUsage

    mensagem = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    pede_secao = plano(sql="", intent="unknown", reason=f"{PEDIDO_DE_SECAO}: prescricao")
    recusado = plano(sql="DELETE FROM cddd.pdvs", usage=AIUsage(model="stub", request={"contexto_completo": True}))
    provider = ScriptedAIProvider(
        [pede_secao, recusado, plano(preenchimento=PREENCHIMENTO)], [resposta("foram 47 unidades")]
    )

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO, RESULTADO]))

    correcao = provider.plan_requests[2]
    assert correcao.error_note
    assert correcao.full_context is True


def test_planilha_pendente_vencida_nao_e_herdada(conversa, catalogo):
    primeira = _pergunta_com_anexo(conversa, Message.Anexo.PLANILHA, PLANILHA, "redes.xlsx")
    _responder(primeira, catalogo, ScriptedAIProvider([_pede_detalhe()]))
    deposito.descartar(primeira.anexo_token)
    provider = ScriptedAIProvider([plano()], [resposta("foram 47 unidades")])

    _responder(_seguimento(conversa), catalogo, provider, FakeQueryExecutor([RESULTADO]))

    assert provider.plan_requests[0].planilha == ""
