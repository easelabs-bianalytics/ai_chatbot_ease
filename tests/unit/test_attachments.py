"""Anexos (ADR-0024): o que sobe ao modelo, o que é recusado e como a
planilha é preenchida.

Cada limite aqui existe por custo ou por memória. Os testes fixam os
números, para ninguém afrouxar um deles sem perceber que mexeu na conta.
"""

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from attachments import deposito
from attachments.imagem import preparar
from attachments.limites import (
    LINHAS_DE_AMOSTRA,
    MAX_CARACTERES_DO_RESUMO,
    MAX_LADO_DA_IMAGEM,
    MAX_LINHAS,
    AnexoRecusado,
)
from attachments.planilha import PedidoDePreenchimento, ler_estrutura, preencher


def _xlsx(linhas) -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = "Rede"
    for linha in linhas:
        aba.append(list(linha))
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def _png(largura, altura, modo="RGB", cor=(200, 30, 30)) -> bytes:
    buffer = io.BytesIO()
    Image.new(modo, (largura, altura), cor).save(buffer, format="PNG")
    return buffer.getvalue()


PEDIDO = PedidoDePreenchimento(
    coluna_chave="Rede",
    chave_no_resultado="rede",
    colunas=(("Sell-out ago/26", "unidades"),),
)


# ------------------------------------------------------------ leitura


def test_so_a_forma_da_planilha_sobe_ao_modelo():
    """Mil linhas de dado, e o resumo leva só cabeçalho, tipo e amostra. Se o
    conteúdo inteiro subisse, uma planilha de 20 mil linhas custaria US$ 2 a
    4 por pergunta."""
    linhas = [("Rede", "UF")] + [(f"Rede secreta {i}", "SP") for i in range(1000)]
    estrutura = ler_estrutura("redes.xlsx", _xlsx(linhas))

    assert estrutura.linhas == 1000
    assert [c.nome for c in estrutura.colunas] == ["Rede", "UF"]
    assert "Rede secreta 999" not in estrutura.resumo
    assert len(estrutura.colunas[0].exemplos) <= LINHAS_DE_AMOSTRA


def test_resumo_tem_teto_de_caracteres():
    cabecalho = tuple(f"Coluna com nome bem comprido número {i}" for i in range(60))
    estrutura = ler_estrutura("larga.xlsx", _xlsx([cabecalho, tuple("x" * 80 for _ in range(60))]))

    assert len(estrutura.resumo) <= MAX_CARACTERES_DO_RESUMO + 40


def test_tipos_das_colunas_sao_inferidos():
    estrutura = ler_estrutura("t.xlsx", _xlsx([("Rede", "Unidades"), ("Pague Menos", 10), ("Drogasil", 12)]))

    assert {c.nome: c.tipo for c in estrutura.colunas} == {"Rede": "texto", "Unidades": "número"}


def test_csv_em_latin1_com_ponto_e_virgula_e_lido():
    """Exportação do Excel no Windows: latin-1 e ';'. Recusar por causa de um
    'ç' seria absurdo."""
    dados = "Rede;Região\nFarmácia São João;Sul\n".encode("latin-1")
    estrutura = ler_estrutura("redes.csv", dados)

    assert [c.nome for c in estrutura.colunas] == ["Rede", "Região"]
    assert estrutura.linhas == 1


def test_planilha_acima_do_limite_de_linhas_e_recusada():
    """Recusa em vez de truncar: preencher metade sem avisar é pior."""
    dados = ("Rede\n" + "\n".join(f"r{i}" for i in range(MAX_LINHAS + 5))).encode()

    with pytest.raises(AnexoRecusado, match="recorte"):
        ler_estrutura("grande.csv", dados)


@pytest.mark.parametrize("nome", ["planilha.xls", "dados.json", "sem_extensao"])
def test_formato_nao_aceito_e_recusado(nome):
    with pytest.raises(AnexoRecusado, match="Formato"):
        ler_estrutura(nome, b"qualquer coisa")


def test_planilha_vazia_e_recusada():
    with pytest.raises(AnexoRecusado, match="vazia"):
        ler_estrutura("v.csv", b"\n\n")


def test_arquivo_corrompido_vira_recusa_legivel():
    with pytest.raises(AnexoRecusado, match="abrir"):
        ler_estrutura("x.xlsx", b"isto nao e um zip")


# -------------------------------------------------------- preenchimento


def test_preenche_casando_pela_chave_e_mantem_o_resto():
    original = _xlsx([
        ("Rede", "Gerente", "Observação"),
        ("Pague Menos", "Ana", "manter"),
        ("drogasil ", "Bia", "manter"),
        ("Rede Nova", "Caio", "manter"),
    ])
    resultado = preencher(
        "redes.xlsx", original, PEDIDO,
        ("rede", "unidades"),
        [("PAGUE MENOS", Decimal("120.5")), ("Drogasil", 80)],
    )

    aba = load_workbook(io.BytesIO(resultado.dados)).active
    linhas = list(aba.iter_rows(values_only=True))
    assert linhas[0] == ("Rede", "Gerente", "Observação", "Sell-out ago/26")
    assert linhas[1] == ("Pague Menos", "Ana", "manter", 120.5)
    assert linhas[2] == ("drogasil ", "Bia", "manter", 80)
    assert linhas[3][3] is None
    assert (resultado.preenchidas, resultado.sem_correspondencia, resultado.total) == (2, 1, 3)


def test_chave_repetida_no_resultado_fica_em_branco():
    """Duas linhas para a mesma rede: escolher uma em silêncio seria
    inventar número."""
    original = _xlsx([("Rede",), ("Pague Menos",)])
    resultado = preencher(
        "r.xlsx", original, PEDIDO, ("rede", "unidades"),
        [("Pague Menos", 10), ("pague menos", 20)],
    )

    aba = load_workbook(io.BytesIO(resultado.dados)).active
    assert aba.cell(row=2, column=2).value is None
    assert resultado.ambiguas == 1
    assert resultado.preenchidas == 0


def test_texto_do_banco_nunca_vira_formula():
    """Valor que começa com '=' executaria no Excel de quem abrir."""
    original = _xlsx([("Rede",), ("A",)])
    resultado = preencher(
        "r.xlsx", original, PEDIDO, ("rede", "unidades"), [("A", "=HYPERLINK(\"x\")")]
    )

    celula = load_workbook(io.BytesIO(resultado.dados)).active.cell(row=2, column=2)
    assert celula.data_type == "s"
    assert celula.value.startswith("'=")


def test_coluna_de_destino_existente_e_reaproveitada():
    original = _xlsx([("Rede", "Sell-out ago/26"), ("A", None)])
    resultado = preencher("r.xlsx", original, PEDIDO, ("rede", "unidades"), [("A", 5)])

    aba = load_workbook(io.BytesIO(resultado.dados)).active
    assert [c.value for c in aba[1]] == ["Rede", "Sell-out ago/26"]
    assert aba.cell(row=2, column=2).value == 5


def test_chave_ausente_na_planilha_e_recusada_com_o_nome():
    with pytest.raises(AnexoRecusado, match="Rede"):
        preencher("r.xlsx", _xlsx([("Produto",), ("X",)]), PEDIDO, ("rede", "unidades"), [])


def test_resultado_sem_as_colunas_pedidas_e_recusado():
    with pytest.raises(AnexoRecusado, match="colunas"):
        preencher("r.xlsx", _xlsx([("Rede",), ("A",)]), PEDIDO, ("outra", "coisa"), [("A", 1)])


def test_csv_sai_preenchido_em_csv():
    original = "Rede;UF\nPague Menos;CE\nOutra;SP\n".encode("utf-8")
    resultado = preencher("r.csv", original, PEDIDO, ("rede", "unidades"), [("pague menos", 7)])

    texto = resultado.dados.decode("utf-8-sig").splitlines()
    assert texto[0] == "Rede;UF;Sell-out ago/26"
    assert texto[1] == "Pague Menos;CE;7"
    assert resultado.sem_correspondencia == 1


# --------------------------------------------------------------- imagem


def test_print_grande_e_reduzido_mantendo_a_proporcao():
    """O custo da imagem é a área. 1920×1080 mediu 2.461 tokens; reduzido,
    cabe em ~1.100."""
    preparada = preparar(_png(1920, 1080))

    assert max(preparada.largura, preparada.altura) == MAX_LADO_DA_IMAGEM
    assert preparada.largura / preparada.altura == pytest.approx(1920 / 1080, rel=0.01)
    assert preparada.reduzida
    assert preparada.tokens_estimados < 1300


def test_imagem_pequena_nao_e_ampliada():
    preparada = preparar(_png(300, 200))

    assert (preparada.largura, preparada.altura) == (300, 200)
    assert not preparada.reduzida


def test_extensao_nao_engana_a_validacao():
    """Um executável renomeado para .png continua não sendo imagem."""
    with pytest.raises(AnexoRecusado):
        preparar(b"MZ\x90\x00 isto e um executavel")


def test_gif_e_recusado():
    buffer = io.BytesIO()
    Image.new("P", (10, 10)).save(buffer, format="GIF")

    with pytest.raises(AnexoRecusado, match="Formato"):
        preparar(buffer.getvalue())


def test_transparencia_vira_fundo_branco():
    """Print com fundo transparente viraria texto preto em fundo preto."""
    preparada = preparar(_png(20, 20, modo="RGBA", cor=(0, 0, 0, 0)))

    pixel = Image.open(io.BytesIO(preparada.dados)).convert("RGB").getpixel((5, 5))
    assert pixel == (255, 255, 255)


# -------------------------------------------------------------- depósito


def test_deposito_guarda_devolve_e_descarta():
    token = deposito.guardar(b"bytes")

    assert deposito.buscar(token) == b"bytes"
    deposito.descartar(token)
    assert deposito.buscar(token) is None


def test_tokens_sao_imprevisiveis():
    """O token trafega até o navegador; se fosse adivinhável, uma pessoa
    baixaria a planilha de outra."""
    tokens = {deposito.guardar(b"x") for _ in range(20)}

    assert len(tokens) == 20
    assert all(len(t) == 32 for t in tokens)


def test_token_vazio_ou_desconhecido_nao_quebra():
    assert deposito.buscar("") is None
    assert deposito.buscar("nao-existe") is None


# ------------------------------------------------- casamento aproximado
# Achado no teste real de 2026-09-21: a planilha dizia "Drogasil" e
# "Panvel", o banco "RAIA DROGASIL" e "PANVEL FARMACIAS" — só 1 de 5 linhas
# casava. Custo zero de token: é regra nossa, não pergunta ao modelo.

REDES_DO_BANCO = [
    ("RAIA DROGASIL", 3851), ("PANVEL FARMACIAS", 462), ("PAGUE MENOS", 777),
    ("DROGARIA DPSP", 632), ("DROGARIA CRISTINA", 17), ("DROGARIA MODERNA", 11),
]


def _preencher_redes(*redes):
    return preencher(
        "r.xlsx", _xlsx([("Rede",), *[(r,) for r in redes]]), PEDIDO, ("rede", "unidades"), REDES_DO_BANCO
    )


def test_nome_contido_em_uma_unica_rede_casa_e_fica_listado():
    resultado = _preencher_redes("Drogasil", "Panvel")

    aba = load_workbook(io.BytesIO(resultado.dados)).active
    assert [aba.cell(row=i, column=2).value for i in (2, 3)] == [3851, 462]
    assert resultado.aproximadas == (("Drogasil", "RAIA DROGASIL"), ("Panvel", "PANVEL FARMACIAS"))


def test_nome_contido_em_varias_redes_nao_casa():
    """"Drogaria" está em três redes: escolher uma seria inventar o número."""
    resultado = _preencher_redes("Drogaria")

    assert resultado.preenchidas == 0
    assert resultado.sem_correspondencia_nomes == ("Drogaria",)


def test_nome_da_planilha_que_contem_o_do_banco_tambem_casa():
    resultado = _preencher_redes("Pague Menos S.A.")

    assert resultado.preenchidas == 1


def test_nome_curto_demais_nao_aproxima():
    """"SA" ou "RJ" casariam com meio banco."""
    resultado = preencher(
        "r.xlsx", _xlsx([("Rede",), ("SA",)]), PEDIDO, ("rede", "unidades"), [("DROGARIA SA RIO", 1)]
    )

    assert resultado.preenchidas == 0


def test_sem_correspondencia_vem_com_os_nomes():
    resultado = _preencher_redes("Drogaria São Paulo", "Rede Que Não Existe", "PAGUE MENOS")

    assert resultado.sem_correspondencia_nomes == ("Drogaria São Paulo", "Rede Que Não Existe")
    assert resultado.aproximadas == ()
    assert resultado.preenchidas == 1
