"""Planilha com várias abas (ADR-0024).

O incidente: em 2026-09-22 o Rubens anexou uma pasta de trabalho com três
abas e pediu a Planilha1. Só a primeira aba era lida, e o Jarvis respondeu
que "consigo ver apenas a estrutura da Planilha3 — reenvie o arquivo com a
Planilha1 acessível". O arquivo estava certo; quem só enxergava uma aba era
o leitor.

O que estes testes fixam:

1. o resumo mostra TODAS as abas, com as colunas de cada uma;
2. o preenchimento escreve na aba que o plano indicou, e só nela;
3. o nome da aba casa por aproximação ("Planilha 3" = "Planilha3"), porque
   quem escreve a pergunta não copia o nome exato;
4. aba inexistente falha dizendo quais existem, em vez de preencher a errada.
"""

import io

import pytest
from openpyxl import Workbook, load_workbook

from attachments.limites import AnexoRecusado
from attachments.planilha import (
    PedidoDePreenchimento,
    ler_estrutura,
    linhas_das_chaves,
    preencher,
)

COLUNAS = ("rede", "unidades")
LINHAS = [("Pague Menos", 47.0), ("Drogasil", 30.0), ("Panvel", 12.0)]


def _pasta_de_trabalho(abas: dict) -> bytes:
    """abas: {nome: [linhas]}, na ordem de inserção."""
    livro = Workbook()
    livro.remove(livro.active)
    for nome, linhas in abas.items():
        aba = livro.create_sheet(nome)
        for linha in linhas:
            aba.append(list(linha))
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


TRES_ABAS = {
    "Planilha1": [("Rede", "Gerente"), ("Pague Menos", "Ana"), ("Drogasil", "Bia")],
    "Planilha2": [("Produto", "Dose"), ("Extrato", "30 ml")],
    "Planilha3": [("Rede", "Unidades"), ("Panvel", None)],
}


def _pedido(aba: str = "") -> PedidoDePreenchimento:
    return PedidoDePreenchimento(
        coluna_chave="Rede",
        chave_no_resultado="rede",
        colunas=(("Unidades", "unidades"),),
        aba=aba,
    )


def test_resumo_traz_todas_as_abas():
    estrutura = ler_estrutura("pasta.xlsx", _pasta_de_trabalho(TRES_ABAS))

    assert estrutura.abas == ("Planilha1", "Planilha2", "Planilha3")
    resumo = estrutura.resumo
    for nome in ("Planilha1", "Planilha2", "Planilha3"):
        assert f'## Aba "{nome}"' in resumo
    # As colunas de uma aba que não é a primeira precisam estar lá: era
    # exatamente isso que faltava quando o modelo disse que não as via.
    assert "Gerente" in resumo and "Produto" in resumo and "Unidades" in resumo


def test_resumo_de_uma_aba_so_nao_fala_de_abas():
    estrutura = ler_estrutura("simples.xlsx", _pasta_de_trabalho({"Dados": TRES_ABAS["Planilha1"]}))

    assert "## Aba" not in estrutura.resumo
    assert estrutura.colunas[0].nome == "Rede"


def test_preenche_a_aba_que_o_plano_indicou():
    dados = _pasta_de_trabalho(TRES_ABAS)

    saida = preencher("pasta.xlsx", dados, _pedido("Planilha3"), COLUNAS, LINHAS)

    livro = load_workbook(io.BytesIO(saida.dados))
    assert [c.value for c in livro["Planilha3"][1]] == ["Rede", "Unidades"]
    assert livro["Planilha3"]["B2"].value == 12.0
    # A Planilha1 tem "Rede" e casaria com o resultado: continua intacta.
    assert [c.value for c in livro["Planilha1"][1]] == ["Rede", "Gerente"]
    assert livro["Planilha1"].max_column == 2


def test_nome_da_aba_casa_sem_exigir_o_espaco():
    dados = _pasta_de_trabalho(TRES_ABAS)

    saida = preencher("pasta.xlsx", dados, _pedido("Planilha 3"), COLUNAS, LINHAS)

    livro = load_workbook(io.BytesIO(saida.dados))
    assert livro["Planilha3"]["B2"].value == 12.0


def test_aba_inexistente_diz_quais_existem():
    dados = _pasta_de_trabalho(TRES_ABAS)

    with pytest.raises(AnexoRecusado) as erro:
        preencher("pasta.xlsx", dados, _pedido("Resumo"), COLUNAS, LINHAS)

    assert "Planilha1" in str(erro.value)


def test_sem_aba_no_plano_preenche_a_primeira():
    dados = _pasta_de_trabalho(TRES_ABAS)

    saida = preencher("pasta.xlsx", dados, _pedido(), COLUNAS, LINHAS)

    livro = load_workbook(io.BytesIO(saida.dados))
    assert livro["Planilha1"]["C2"].value == 47.0


def test_a_resposta_fala_das_linhas_da_aba_preenchida():
    """A tabela na tela sai da mesma aba em que o arquivo foi escrito."""
    dados = _pasta_de_trabalho(TRES_ABAS)

    escolhidas = linhas_das_chaves("pasta.xlsx", dados, _pedido("Planilha3"), COLUNAS, LINHAS)

    assert escolhidas == [("Panvel", 12.0)]
