"""Preenchimento de várias colunas e a tabela transcrita de um print (ADR-0024).

A planilha real de 2026-09-21 pedia PX Ease, PX Aché, variação, unidades e
market share lado a lado: uma chave, cinco colunas. E o mesmo pedido chegou
como print de uma tabela vazia.
"""

import io

import pytest
from openpyxl import load_workbook

from attachments.limites import AnexoRecusado
from attachments.planilha import (
    MAX_LINHAS_DO_PRINT,
    PedidoDePreenchimento,
    ler_estrutura,
    montar_de_tabela,
    preencher,
)
from tests.unit.test_attachments import _xlsx

PEDIDO = PedidoDePreenchimento(
    coluna_chave="REPRESENTANTE",
    chave_no_resultado="representante",
    colunas=(("PX EASE", "px_ease"), ("VAR %", "var_pct"), ("MKT SHARE", "share")),
)
COLUNAS = ("representante", "px_ease", "var_pct", "share")
LINHAS = [("MARTA ELOISA", 120, 0.15, 0.31), ("JONATHAN ABRAHAO", 90, -0.02, 0.22)]


def test_preenche_varias_colunas_de_uma_vez():
    original = _xlsx([("REPRESENTANTE", "PX EASE", "VAR %"), ("MARTA ELOISA", None, None), ("JONATHAN ABRAHAO",)])

    resultado = preencher("r.xlsx", original, PEDIDO, COLUNAS, LINHAS)

    linhas = list(load_workbook(io.BytesIO(resultado.dados)).active.iter_rows(values_only=True))
    # MKT SHARE não existia: é criada no fim, sem mexer na ordem das outras.
    assert linhas[0] == ("REPRESENTANTE", "PX EASE", "VAR %", "MKT SHARE")
    assert linhas[1] == ("MARTA ELOISA", 120, 0.15, 0.31)
    assert linhas[2] == ("JONATHAN ABRAHAO", 90, -0.02, 0.22)
    assert resultado.preenchidas == 2


def test_varias_colunas_em_csv():
    original = "REPRESENTANTE;PX EASE\nMARTA ELOISA;\n".encode("utf-8")

    resultado = preencher("r.csv", original, PEDIDO, COLUNAS, LINHAS)

    texto = resultado.dados.decode("utf-8-sig").splitlines()
    assert texto[0] == "REPRESENTANTE;PX EASE;VAR %;MKT SHARE"
    assert texto[1] == "MARTA ELOISA;120;0.15;0.31"


def test_coluna_do_resultado_faltando_recusa_o_preenchimento_inteiro():
    """Se uma das colunas não veio, não dá para saber o que mais está errado."""
    pedido = PedidoDePreenchimento("REPRESENTANTE", "representante", (("PX", "px_ease"), ("X", "nao_existe")))

    with pytest.raises(AnexoRecusado, match="colunas"):
        preencher("r.xlsx", _xlsx([("REPRESENTANTE",), ("MARTA ELOISA",)]), pedido, COLUNAS, LINHAS)


def test_pedido_sai_do_formato_do_plano():
    pedido = PedidoDePreenchimento.do_plano({
        "coluna_chave": "REPRESENTANTE",
        "chave_no_resultado": "representante",
        "colunas": [{"coluna_destino": "PX EASE", "valor_no_resultado": "px_ease"}],
    })

    assert pedido.colunas == (("PX EASE", "px_ease"),)


def test_tabela_do_print_vira_planilha_em_memoria():
    dados = montar_de_tabela(
        ["REPRESENTANTE", "PX EASE YTD 2026", ""],
        [["MARTA ELOISA", ""], ["JONATHAN ABRAHAO", ""], ["", ""]],
    )

    estrutura = ler_estrutura("tabela_do_print.xlsx", dados)
    assert [c.nome for c in estrutura.colunas] == ["REPRESENTANTE", "PX EASE YTD 2026"]
    assert estrutura.linhas == 2  # a linha vazia do print não vira linha


def test_print_sem_cabecalho_nao_vira_planilha():
    with pytest.raises(AnexoRecusado, match="tabela"):
        montar_de_tabela(["", "  "], [["a"]])


def test_print_com_linhas_demais_e_cortado():
    """Um print legível não tem 300 linhas; o teto só segura erro do modelo."""
    dados = montar_de_tabela(["NOME"], [[f"n{i}"] for i in range(MAX_LINHAS_DO_PRINT + 50)])

    assert ler_estrutura("t.xlsx", dados).linhas == MAX_LINHAS_DO_PRINT


def test_coluna_de_texto_mostra_a_amostra_inteira_no_resumo():
    """Com dois exemplos, o modelo filtrou por dois dos três representantes."""
    dados = montar_de_tabela(
        ["REPRESENTANTE", "PX"], [["MARTA ELOISA", 1], ["JONATHAN ABRAHAO", 2], ["ALEXANDRE CIMINI", 3]]
    )

    resumo = ler_estrutura("t.xlsx", dados).resumo

    assert "ALEXANDRE CIMINI" in resumo


# ---------------------------------- a resposta fala só do que foi pedido
# Produção, 2026-09-22: a planilha tinha 3 representantes, a consulta trouxe
# 34 e a resposta falou dos 34. O preenchimento acertava; a leitura, não.

RESULTADO_COMPLETO = (
    ("representante", "px_ease"),
    [("MARTA ELOISA", 120), ("DANILO SOUZA", 90), ("HERMES BIZZOTO", 300), ("CAROLINE SILVA", 70)],
)
PEDIDO_DE_UMA = PedidoDePreenchimento(
    coluna_chave="REPRESENTANTE", chave_no_resultado="representante", colunas=(("PX", "px_ease"),)
)


def test_so_as_linhas_da_planilha_entram_na_resposta():
    from attachments.planilha import linhas_das_chaves

    planilha = _xlsx([("REPRESENTANTE",), ("MARTA ELOISA",), ("HERMES",)])

    escolhidas = linhas_das_chaves("r.xlsx", planilha, PEDIDO_DE_UMA, *RESULTADO_COMPLETO)

    # Na ordem da planilha; "HERMES" casa com "HERMES BIZZOTO" por estar
    # contido em uma única chave. Grafia trocada dentro da palavra
    # ("BIZOTTO" × "BIZZOTO") é trabalho da consulta, com ILIKE por pedaço.
    assert escolhidas == [("MARTA ELOISA", 120), ("HERMES BIZZOTO", 300)]


def test_chave_que_nao_casa_nao_entra():
    from attachments.planilha import linhas_das_chaves

    planilha = _xlsx([("REPRESENTANTE",), ("QUEM NAO EXISTE",)])

    assert linhas_das_chaves("r.xlsx", planilha, PEDIDO_DE_UMA, *RESULTADO_COMPLETO) == []


def test_resumo_avisa_quando_a_amostra_cobre_a_planilha():
    dados = montar_de_tabela(["REDE"], [["RAIA DROGASIL"], ["PANVEL"], ["PAGUE MENOS"]])

    assert "todas elas" in ler_estrutura("r.xlsx", dados).resumo


def test_resumo_nao_promete_cobertura_em_planilha_grande():
    dados = montar_de_tabela(["REDE"], [[f"rede {i}"] for i in range(40)])

    assert "todas elas" not in ler_estrutura("r.xlsx", dados).resumo
