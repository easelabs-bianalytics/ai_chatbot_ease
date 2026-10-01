"""Planilha no orquestrador, Jarvis 3.0 (ADR-0031): as conversas 37 e 44.

Conversa 44 (Paulo, 2026-09-29): 230 médicos com CRM, pedido de trazer o
território da UTC de cada um. O Jarvis varreu `audit.medico` inteiro, a
consulta parou em 50 mil linhas, 84 de 230 foram preenchidos, a resposta
disse "não é seguro preencher" acima de "Preenchi 84 de 230", e na correção
seguinte pediu o arquivo de novo. Estes testes fixam o contrário:

1. a consulta parte de `anexo.<aba>`, com os CRMs como parâmetro;
2. o preenchimento volta pela `_linha`, e roda uma vez só;
3. a redação recebe o que a planilha ganhou, antes de escrever;
4. a planilha continua na conversa para a correção seguinte;
5. nada sai se a conferência achar alteração fora do pedido, nem pela metade;
6. relatório e análise sobre a própria planilha.
"""

import io

import pytest
from openpyxl import Workbook, load_workbook

from ai_orchestrator import canned
from ai_orchestrator import orchestrator as orquestrador
from ai_orchestrator.models import AIReply
from ai_orchestrator.orchestrator import handle_message
from attachments import deposito, qa
from attachments.planilha import NOME_DAS_NOTAS, ler_estrutura
from catalog.loader import load_catalog
from conversations.models import Conversation
from datasource.executors.fake import FakeQueryExecutor, make_result
from messaging.channels.fake import FakeChannel
from messaging.models import Message
from tests.fakes.providers import ScriptedAIProvider, plano, resposta

pytestmark = pytest.mark.django_db


def _xlsx(linhas, titulo="Painel") -> bytes:
    livro = Workbook()
    livro.active.title = titulo
    for linha in linhas:
        livro.active.append(list(linha))
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


PAINEL = _xlsx([
    ("Painel Médico - Simone (setor 3000)",),
    ("Nome do médico", "CRM", "UF de atendimento"),
    ("BRENDON DUTRA", "MG0104608", "BA"),
    ("ELISA OLIVEIRA", "MG0049899", "MG"),
    ("FARES NETO", "PR0023511", "PR"),
])

SQL_DO_ANEXO = (
    "SELECT a._linha, btrim(fv.desc_territorio) AS representante\n"
    "FROM anexo.painel a\n"
    "LEFT JOIN audit.medico m ON upper(m.crm) = upper(a.crm)\n"
    "LEFT JOIN cddd.forca_vendas fv ON fv.cod_utc = m.utc_codigo"
)
PREENCHIMENTO = {
    "coluna_chave": "CRM",
    "chave_no_resultado": "_linha",
    "colunas": [{"coluna_destino": "Representante", "valor_no_resultado": "representante",
                 "justificativa": "território da UTC do endereço do médico"}],
    "aba": "",
}
RESULTADO = make_result(("_linha", "representante"), [(3, "SEM REP"), (4, "HERMES"), (5, None)])


@pytest.fixture(scope="module")
def catalogo():
    return load_catalog()


@pytest.fixture
def conversa(django_user_model):
    user = django_user_model.objects.create_user("paulo", password="x")
    return Conversation.objects.create(user=user)


def _com_planilha(conversa, texto="Adicione o representante da UTC de cada médico", dados=PAINEL,
                  nome="Painel Médico.xlsx", cid="c-1"):
    return Message.objects.create(
        conversation=conversa, direction=Message.Direction.INBOUND, content=texto, client_message_id=cid,
        status=Message.Status.RECEIVED, anexo_tipo=Message.Anexo.PLANILHA, anexo_nome=nome,
        anexo_resumo=ler_estrutura(nome, dados).resumo, anexo_token=deposito.guardar(dados),
    )


def _seguimento(conversa, texto, cid):
    return Message.objects.create(
        conversation=conversa, direction=Message.Direction.INBOUND, content=texto, client_message_id=cid,
    )


def _responder(mensagem, catalogo, provider, executor):
    return handle_message(mensagem, channel=FakeChannel(), provider=provider, executor=executor, catalog=catalogo)


def _devolvida(mensagem):
    mensagem.refresh_from_db()
    return load_workbook(io.BytesIO(deposito.buscar(mensagem.anexo_resposta_token)))


def _preencher(conversa, catalogo, **campos):
    mensagem = _com_planilha(conversa)
    executor = FakeQueryExecutor([RESULTADO])
    provider = ScriptedAIProvider(
        [plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO, operacao_da_planilha="enriquecer", **campos)],
        [resposta("Um dos três médicos está com HERMES.")],
    )
    reply = _responder(mensagem, catalogo, provider, executor)
    return mensagem, reply, provider, executor


def test_consulta_parte_da_planilha_com_os_crms_como_parametro(conversa, catalogo):
    _, _, provider, executor = _preencher(conversa, catalogo)

    assert "Tabela SQL: anexo.painel" in provider.plan_requests[0].planilha
    sql = executor.executed[0]
    assert "anexo__painel(_linha" in sql and "FROM anexo__painel a" in sql
    assert "MG0104608" not in sql
    assert executor.params[0][1] == ["MG0104608", "MG0049899", "PR0023511"]


def test_preenchimento_volta_pela_linha_e_roda_uma_vez(conversa, catalogo):
    mensagem, reply, _, executor = _preencher(conversa, catalogo)

    aba = _devolvida(mensagem)["Painel"]
    assert [aba.cell(row=n, column=4).value for n in (2, 3, 4, 5)] == ["Representante", "SEM REP", "HERMES", None]
    assert len(executor.executed) == 1
    assert "Preenchi **2 de 3 linhas**" in reply.reply_text
    assert "sem dado no banco:** PR0023511" in reply.reply_text


def test_redacao_recebe_a_planilha_devolvida_antes_de_escrever(conversa, catalogo):
    """A resposta e a nota do sistema falam a mesma coisa: a redação recebe
    o relatório da planilha já alterada e conferida."""
    _, reply, provider, _ = _preencher(conversa, catalogo)

    relatorio = provider.answer_requests[0].planilha_devolvida
    assert "2 de 3 linhas ganharam valor" in relatorio
    assert "nenhuma célula original mudou" in relatorio
    assert reply.decision == AIReply.Decision.ANSWERED
    assert reply.raw_response["planilha_alterada"]["conferencia"] == []


def test_arquivo_devolvido_traz_as_notas_do_jarvis(conversa, catalogo):
    mensagem, reply, _, _ = _preencher(conversa, catalogo)

    livro = _devolvida(mensagem)
    assert livro.sheetnames == ["Painel", NOME_DAS_NOTAS]
    notas = list(livro[NOME_DAS_NOTAS].iter_rows(values_only=True))
    assert notas[1][0] == "Aba Painel · coluna Representante"
    assert notas[1][2] == "território da UTC do endereço do médico"
    assert "Notas do Jarvis" in reply.reply_text


def test_valor_que_domina_a_coluna_e_dito(conversa, catalogo):
    """"SEM REP" em quase tudo é o achado da planilha, e tem de aparecer."""
    mensagem = _com_planilha(conversa)
    todos = make_result(("_linha", "representante"), [(3, "SEM REP"), (4, "SEM REP"), (5, "SEM REP")])
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [resposta("ok")])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([todos]))

    assert "**SEM REP** aparece em 3 das 3 linhas preenchidas de `Representante`" in reply.reply_text


def test_correcao_seguinte_ainda_tem_a_planilha(conversa, catalogo):
    """"Você fez errado, use só os dados da planilha" — a planilha continua
    na conversa, com `anexo.painel` liberado para a nova consulta."""
    _preencher(conversa, catalogo)
    correcao = _seguimento(conversa, "Você fez errado, faça o join só com os CRMs da planilha", "c-2")
    executor = FakeQueryExecutor([RESULTADO])
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [resposta("ok")])

    reply = _responder(correcao, catalogo, provider, executor)

    assert "anexo.painel" in provider.plan_requests[0].planilha
    assert reply.decision == AIReply.Decision.ANSWERED
    assert executor.params[0] is not None
    assert "Preenchi **2 de 3 linhas**" in reply.reply_text


def test_resposta_que_nao_usou_a_planilha_encerra_a_heranca(conversa, catalogo):
    _preencher(conversa, catalogo)
    outra = _seguimento(conversa, "quanto vendemos ontem?", "c-2")
    _responder(outra, catalogo, ScriptedAIProvider([plano()], [resposta("foram 47 unidades")]),
               FakeQueryExecutor([make_result(("und",), [(47,)])]))
    terceira = _seguimento(conversa, "e na semana?", "c-3")
    provider = ScriptedAIProvider([plano()], [resposta("foram 47 unidades")])

    _responder(terceira, catalogo, provider, FakeQueryExecutor([make_result(("und",), [(47,)])]))

    assert provider.plan_requests[0].planilha == ""


def test_conferencia_que_falha_nao_entrega_o_arquivo(conversa, catalogo, monkeypatch):
    falha = qa.Conferencia(problemas=['2 células da aba "Painel" mudaram fora do pedido (linha 3, coluna 2)'])
    monkeypatch.setattr(orquestrador.conferencia_da_planilha, "conferir", lambda *a, **k: falha)

    mensagem, reply, _, _ = _preencher(conversa, catalogo)

    mensagem.refresh_from_db()
    assert mensagem.anexo_resposta_token == ""
    assert "Não entreguei a planilha" in reply.reply_text
    assert "Preenchi" not in reply.reply_text


def test_consulta_cortada_nao_vira_planilha_pela_metade(conversa, catalogo):
    mensagem = _com_planilha(conversa)
    cortado = make_result(RESULTADO.columns, RESULTADO.rows, truncated=True)
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [resposta("ok")])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([cortado, cortado]))

    mensagem.refresh_from_db()
    assert mensagem.anexo_resposta_token == ""
    assert canned.PLANILHA_CORTADA in reply.reply_text
    assert "Preenchi" not in reply.reply_text


def test_analise_sobre_a_planilha_nao_pede_casamento(conversa, catalogo):
    """"Quantos médicos por UF?" é análise da planilha: resposta, sem
    arquivo e sem o aviso de "não sei qual coluna casa"."""
    mensagem = _com_planilha(conversa, texto="quantos médicos por UF?")
    por_uf = make_result(("uf", "medicos"), [("BA", 1), ("MG", 1), ("PR", 1)])
    provider = ScriptedAIProvider(
        [plano("SELECT a.uf_de_atendimento AS uf, count(*) AS medicos FROM anexo.painel a GROUP BY 1",
               operacao_da_planilha="analisar")],
        [resposta("Um médico em cada UF.")],
    )

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([por_uf]))

    assert canned.PLANILHA_SEM_CASAMENTO not in reply.reply_text
    assert reply.raw_response["anexo_sql"] == ["painel"]
    assert reply.raw_response["operacao_da_planilha"] == "analisar"


def test_relatorio_vira_abas_novas_com_grafico(conversa, catalogo):
    mensagem = _com_planilha(conversa, texto="monte um resumo por UF e por potencial na planilha")
    por_uf = make_result(("uf", "medicos"), [("BA", 1), ("MG", 1), ("PR", 1)])
    total = make_result(("total",), [(3,)])
    consultas = (
        {"titulo": "Médicos por UF", "sql": "SELECT a.uf_de_atendimento AS uf, count(*) AS medicos FROM anexo.painel a GROUP BY 1",
         "reference_query_id": "", "aba_nova": "Resumo por UF", "grafico_na_aba": "barras"},
        {"titulo": "Total", "sql": "SELECT count(*) AS total FROM anexo.painel a",
         "reference_query_id": "", "aba_nova": "Total"},
    )
    provider = ScriptedAIProvider(
        [plano("", consultas=consultas, operacao_da_planilha="relatorio")],
        [resposta("São três médicos, um por UF.")],
    )

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([por_uf, total]))

    livro = _devolvida(mensagem)
    assert livro.sheetnames == ["Painel", "Resumo por UF", "Total", NOME_DAS_NOTAS]
    assert len(livro["Resumo por UF"]._charts) == 1
    assert "**Abas novas:**" in reply.reply_text
    assert "Aba nova Resumo por UF: 3 linhas, com gráfico." in provider.answer_requests[0].planilha_devolvida


def test_aba_que_nao_existe_volta_para_a_correcao(conversa, catalogo):
    """O validador recusa `anexo.outra` com a lista das abas certas, e é
    isso que vai na correção única."""
    mensagem = _com_planilha(conversa)
    provider = ScriptedAIProvider(
        [plano("SELECT a._linha FROM anexo.outra a"),
         plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)],
        [resposta("ok")],
    )

    _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    assert "anexo.painel" in provider.plan_requests[1].error_note


def test_redacao_sabe_quando_a_planilha_nao_mudou(conversa, catalogo):
    """2026-10-01: "a coluna Potencial foi preenchida" acima de "Não consegui
    preencher a planilha". Quando nada muda, a redação é avisada."""
    mensagem = _com_planilha(conversa)
    errado = {**PREENCHIMENTO, "colunas": [{"coluna_destino": "X", "valor_no_resultado": "nao_existe"}]}
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=errado)], [resposta("ok")])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    assert "NÃO foi alterada" in provider.answer_requests[0].planilha_devolvida
    assert "Não consegui preencher a planilha" in reply.reply_text


def test_cobertura_por_coluna_com_varias_colunas(conversa, catalogo):
    """"Preenchi 513 de 513" com o representante em 493: com várias colunas,
    cada uma diz a sua."""
    mensagem = _com_planilha(conversa)
    duas = {**PREENCHIMENTO, "colunas": [
        {"coluna_destino": "Representante", "valor_no_resultado": "representante"},
        {"coluna_destino": "Critério", "valor_no_resultado": "criterio"},
    ]}
    resultado = make_result(("_linha", "representante", "criterio"),
                            [(3, "ANA", "CRM"), (4, None, "não identificado"), (5, "BIA", "CRM")])
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=duas)], [resposta("ok")])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([resultado]))

    assert "Por coluna: `Representante` 2 de 3 · `Critério` 3 de 3." in reply.reply_text


def test_sem_nenhuma_celula_escrita_nao_ha_arquivo(conversa, catalogo):
    mensagem = _com_planilha(conversa)
    vazio = make_result(("_linha", "representante"), [(3, None), (4, None), (5, None)])
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [resposta("ok")])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([vazio]))

    mensagem.refresh_from_db()
    assert mensagem.anexo_resposta_token == ""
    assert "Não alterei a planilha: a consulta não trouxe valor para nenhuma das 3 linhas pedidas." in reply.reply_text
    assert "NÃO foi alterada" in provider.answer_requests[0].planilha_devolvida


# ------------------------------------------- tabela com a planilha entregue (2026-10-01)
# Conversa 52: a planilha saiu certa e a tela mostrou 100 de 230 linhas de
# `_linha · representante · gr`. Com o arquivo devolvido, a tabela só aparece
# se couber inteira — e com o identificador da planilha no lugar do `_linha`.


def _com_blocos(texto="Preenchi a planilha."):
    from tests.fakes.providers import resposta as r

    return r(texto, blocos=({"tipo": "texto", "texto": texto}, {"tipo": "tabela", "consulta": 0, "colunas": []}))


def test_tabela_que_nao_cabe_sai_quando_a_planilha_foi_entregue(conversa, catalogo):
    linhas = [("Médico %d" % i, "MG%07d" % i) for i in range(150)]
    dados = _xlsx([("Nome do médico", "CRM"), *linhas])
    mensagem = _com_planilha(conversa, dados=dados)
    resultado = make_result(("_linha", "representante"), [(i + 2, "HERMES") for i in range(150)])
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [_com_blocos()])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([resultado]))

    assert [b["tipo"] for b in reply.raw_response["blocos"]] == ["texto"]
    assert "Preenchi **150 de 150 linhas**" in reply.reply_text


def test_tabela_que_cabe_mostra_o_crm_no_lugar_do_numero_da_linha(conversa, catalogo):
    mensagem = _com_planilha(conversa)
    provider = ScriptedAIProvider([plano(SQL_DO_ANEXO, preenchimento=PREENCHIMENTO)], [_com_blocos()])

    reply = _responder(mensagem, catalogo, provider, FakeQueryExecutor([RESULTADO]))

    tabela = next(b for b in reply.raw_response["blocos"] if b["tipo"] == "tabela")
    dados = reply.raw_response["dados_blocos"][str(tabela["consulta"])]
    assert dados["columns"][0] == "CRM"
    assert [linha[0] for linha in dados["rows"]] == ["MG0104608", "MG0049899", "PR0023511"]
    assert "_linha" not in tabela["colunas"]
